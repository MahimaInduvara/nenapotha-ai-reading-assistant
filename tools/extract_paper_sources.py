from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def extract_document(source: Path, destination: Path) -> None:
    document = Document(source)
    lines: list[str] = [f"SOURCE: {source}", "", "PARAGRAPHS"]
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.replace("\t", " ").strip()
        if text:
            style_name = paragraph.style.name if paragraph.style is not None else "(no style)"
            lines.append(f"P{index:04d} [{style_name}] {text}")

    lines.extend(["", "TABLES"])
    for table_index, table in enumerate(document.tables):
        lines.append(f"TABLE {table_index}")
        for row_index, row in enumerate(table.rows):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            lines.append(f"R{row_index:03d} | " + " | ".join(cells))

    lines.extend(["", "SECTIONS"])
    for index, section in enumerate(document.sections):
        lines.append(
            " ".join(
                [
                    f"S{index}",
                    f"width={section.page_width}",
                    f"height={section.page_height}",
                    f"left={section.left_margin}",
                    f"right={section.right_margin}",
                    f"top={section.top_margin}",
                    f"bottom={section.bottom_margin}",
                    f"columns_xml={section._sectPr.xml}",
                ]
            )
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("documents", nargs="+", type=Path)
    args = parser.parse_args()

    for source in args.documents:
        if not source.is_file():
            raise FileNotFoundError(source)
        destination = args.output_dir / f"{source.stem}.txt"
        extract_document(source, destination)
        print(destination)


if __name__ == "__main__":
    main()
