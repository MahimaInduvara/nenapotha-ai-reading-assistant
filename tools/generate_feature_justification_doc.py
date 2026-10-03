from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from generate_thesis_breakdown import (
    BLUE,
    CORAL,
    MID_GREY,
    NAVY,
    PALE_ORANGE,
    PALE_TEAL,
    add_body,
    add_bullets,
    add_callout,
    add_numbered,
    add_table,
    add_toc,
    setup_document,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
DOCX = OUT / "ReadBuddy_AI_Supervisor_Feature_Justification_28277.docx"


def title_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run("READBUDDY AI")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(27)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("Research Justification for Every System Feature")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("Supervisor Discussion and Viva Preparation Document")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(CORAL)

    for text, size, bold in [
        ("Design and Development of an AI-Enhanced Bilingual Reading and Comprehension Assistant for Grade 1 and Grade 2 Primary School Students", 14, True),
        ("On-Device Sinhala Letter Recognition and Sinhala Text-Difficulty Classification", 13, False),
        ("L. M. I. Silva", 14, True),
        ("Student ID: 28277", 12, False),
        ("BSc (Hons) in Software Engineering — Final Year Research Project", 11, False),
        ("NSBM Green University — Faculty of Computing", 11, False),
        ("2026", 11, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(25 if text == "L. M. I. Silva" else 9)
        rr = p.add_run(text)
        rr.font.name = "Times New Roman"
        rr.font.size = Pt(size)
        rr.bold = bold

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    rr = p.add_run("Supervisor: ______________________________________________")
    rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor.from_string(MID_GREY)
    doc.add_page_break()


def front_matter(doc):
    doc.add_heading("Document Purpose", level=1)
    add_body(doc, "This document explains why each ReadBuddy AI feature was included, how it contributes to the research aim and objectives, how it can be evaluated, and what limitation must be disclosed. It is designed for a supervisor discussion, proposal/interim presentation and viva preparation. It distinguishes core research contributions from application-support features so that ordinary product functionality is not incorrectly presented as a separate research contribution.")
    add_callout(
        doc,
        "Central defence statement",
        "The on-device Sinhala letter-recognition CNN and the Grade 1/2 Sinhala text-difficulty classifier are the two core research components. All remaining features deliver the intervention, collect evidence, support accessibility and personalization, or allow responsible adults to interpret progress.",
        PALE_TEAL,
    )
    add_callout(
        doc,
        "Evidence rule",
        "A feature may be described as implemented when its code path works. It may be described as validated only when appropriate evaluation evidence exists. It may be described as educationally effective only after an ethics-approved learning-outcome study. These three claims are not interchangeable.",
        PALE_ORANGE,
    )

    doc.add_heading("Table of Contents", level=1)
    add_toc(doc)
    doc.add_page_break()


def research_structure(doc):
    doc.add_heading("1. Feature-Justification Structure", level=1)
    add_body(doc, "For every feature, use the following six-part structure. This turns a product explanation into a research explanation and ensures that the answer is traceable to evidence.")
    add_numbered(doc, [
        "Need or problem — identify the educational, technical or stakeholder problem addressed by the feature.",
        "Design decision — explain why this method was selected instead of a reasonable alternative.",
        "Research alignment — connect the feature to the research question and one or more objectives.",
        "Technical operationalization — state the input, process, output and decision rule.",
        "Evaluation — identify the metric, test, comparison or user evidence that will determine success.",
        "Limitation — state what the current evidence does not establish and what remains to be completed.",
    ])
    add_callout(doc, "Reusable oral-answer template", "I added [feature] because [specific problem]. I selected [method] because [constraint or comparison]. It supports [objective] by [research contribution]. The feature takes [input], applies [process], and produces [output]. I will evaluate it using [metrics/evidence]. At present, the main limitation is [honest boundary].")

    doc.add_heading("2. Research Aim, Question and Objectives", level=1)
    add_callout(doc, "Research aim", "To design, develop and evaluate an AI-enhanced bilingual mobile reading and comprehension assistant for Grade 1 and Grade 2 students by integrating an on-device Sinhala handwritten-letter CNN and an interpretable Sinhala Grade 1/2 text-difficulty classifier.")
    add_callout(doc, "Primary research question", "How can an AI-enhanced bilingual mobile reading assistant be designed and evaluated to provide offline Sinhala letter-tracing classification for Grade 1 learners and automated Sinhala text-difficulty classification for Grade 1 and Grade 2 reading content?")
    add_table(doc, ["Objective", "Purpose", "Feature relationship"], [
        ("RO1 — Identify", "Identify educational, application and evidence gaps affecting letter-tracing feedback and grade-appropriate content selection.", "Problem definition, teacher requirements, scope and stakeholder needs."),
        ("RO2 — Analyse", "Analyse Sinhala handwriting recognition, low-resource readability, on-device deployment and evaluation strategies.", "Technology selection and design trade-offs."),
        ("RO3 — Develop", "Develop the TFLite CNN, logistic-regression text classifier and integrated Flutter artifact.", "Core models and supporting mobile workflows."),
        ("RO4 — Evaluate", "Evaluate the components, integration, reliability, validity and limitations using transparent evidence.", "Metrics, baselines, tests, usability and field-evidence plan."),
    ])

    doc.add_heading("3. Three-Layer System Structure", level=1)
    add_table(doc, ["Layer", "Included features", "Research function"], [
        ("Core research layer", "Sinhala letter CNN; Grade 1/2 text-difficulty classifier; mobile integration.", "Produces the principal technical contribution and quantitative evaluation."),
        ("Educational delivery layer", "Letters, tracing, pillam, stories, TTS, vocabulary, quizzes and recommendations.", "Provides authentic inputs and a learning workflow in which the models can be demonstrated."),
        ("Evidence and stakeholder layer", "Profiles, session logging, progress, weekly insight, parent/teacher access and adult AI support.", "Collects evidence and makes it interpretable to responsible adults."),
    ])


def core_features(doc):
    doc.add_heading("4. Core Research Feature Justifications", level=1)

    doc.add_heading("4.1 On-Device Sinhala Letter-Recognition CNN", level=2)
    add_body(doc, "Problem addressed: A normal tracing canvas records strokes but cannot objectively determine whether the intended Sinhala letter was produced. Reviewed Sinhala handwritten-character research demonstrates the feasibility of CNN-based recognition [4]–[6], but does not establish a Grade 1 child-facing on-device feedback workflow.")
    add_body(doc, "Reason for inclusion: The CNN converts a child trace into a measurable predicted class and confidence score. TensorFlow Lite was chosen to support offline operation, reduce network delay and avoid uploading every drawing to an external service. On-device inference is also a testable mobile-computing decision rather than merely an interface feature [10].")
    add_table(doc, ["Research element", "Definition"], [
        ("Input", "Captured drawing image resized to 64×64 grayscale and normalized to [0,1]."),
        ("Process", "Local TFLite inference with a 454-way softmax output."),
        ("Output", "Top-1 predicted class and confidence value."),
        ("Current decision rule", "Expected-label match and confidence above 0.50; geometric coverage is a fallback."),
        ("Evaluation", "Validation accuracy, macro-F1, per-class recall, confusion matrix, top-k accuracy, calibration, latency and end-to-end mapping tests."),
        ("RO alignment", "RO2, RO3 and RO4."),
    ])
    add_callout(doc, "Required limitation statement", "Approximately 94% is validation accuracy on reference data, not evidence of equivalent accuracy on children's traces. The current numeric class IDs must be mapped to Sinhala Unicode letters, only five letters are exposed in the tracing interface, and the exact training notebook for the shipped model must be recovered before the pipeline can be described as reproducible.", PALE_ORANGE)

    doc.add_heading("4.2 Grade 1/2 Sinhala Text-Difficulty Classifier", level=2)
    add_body(doc, "Problem addressed: Reading passages can be assigned grade labels manually without a reproducible check of measurable text difficulty. Sinhala also lacks the large expert-labelled readability resources used by many high-resource-language systems [7]–[9].")
    add_body(doc, "Reason for inclusion: A binary logistic-regression model provides an interpretable low-resource baseline. It is lightweight, runs offline in Dart and exposes the contribution of word count, average word length and average sentence length. This is more defensible than adopting a large opaque model for a dataset that is currently very small.")
    add_table(doc, ["Research element", "Definition"], [
        ("Input", "Sinhala story text."),
        ("Features", "Number of words, average word length and average sentence length."),
        ("Process", "Binary logistic regression and a 0.50 probability threshold."),
        ("Output", "Estimated Grade 1 or Grade 2 difficulty."),
        ("Evaluation", "Stratified cross-validation, macro-F1, balanced accuracy, per-class recall, confusion matrix, majority baseline and educator agreement."),
        ("RO alignment", "RO1, RO2, RO3 and RO4."),
    ])
    add_callout(doc, "Required limitation statement", "The model was fitted on only 26 samples (9 Grade 1 and 17 Grade 2), and 69.2% is training accuracy. It is an exploratory proof of concept, not a validated production classifier. A larger balanced, independently educator-labelled corpus is required.", PALE_ORANGE)

    doc.add_heading("4.3 Integrated Mobile Research Artifact", level=2)
    add_body(doc, "Reason for inclusion: The contribution is not limited to training two isolated models. The research investigates how letter-level computer vision and passage-level readability assessment can be combined within one bilingual early-literacy workflow. Flutter delivers the artifact, TFLite performs local image inference, the Dart classifier checks story difficulty, and Firebase persists authorised evidence.")
    add_table(doc, ["Evaluation area", "Evidence required"], [
        ("Functional correctness", "Model load, preprocessing shape, output size, Unicode mapping and UI verdict tests."),
        ("Mobile feasibility", "Model size, median/p95 latency, memory use and crash-free runs on a physical Android device."),
        ("Workflow usability", "Task completion, navigation errors, help requests and educator/user ratings."),
        ("Reliability", "Offline behaviour, failure handling, repeatability and persistence consistency."),
        ("Validity", "Reference-data results separated from child-field results and educational outcomes."),
    ])


def feature_matrices(doc):
    doc.add_heading("5. Educational Delivery Feature Matrix", level=1)
    add_body(doc, "These features support the research artifact. They should not each be presented as independent AI contributions.")
    add_table(doc, ["Feature", "Research justification and role", "Evaluation and honest boundary"], [
        ("Grade 1/2 scope", "Bounds the population and aligns the interface, content and binary classifier with the implemented evidence. It prevents unsupported Grade 3 claims.", "Educator content review and scope consistency audit. Grade 3 is out of scope."),
        ("Sinhala/English interface", "Improves access to instructions in the familiar language while offering controlled English support. It is a delivery and accessibility decision.", "Bilingual task completion, terminology review and navigation errors. Do not claim improved English proficiency without outcome evidence."),
        ("Pre-reading sequence", "Scaffolds listening, sound identification, visual matching and letter recognition before complex reading. A 60% rule operationalises progression.", "Stage accuracy, attempts and completion. The screen currently exists but is not connected to main navigation."),
        ("Letter lessons with TTS and examples", "Links the visual symbol to its sound and an example word rather than presenting isolated characters.", "Letter identification and educator pronunciation/content review. Device TTS is support, not a validated pronunciation model."),
        ("Letter-recognition task", "Measures association among a picture, word and initial letter; produces per-letter errors.", "Accuracy, attempts, time and confusion patterns."),
        ("Letter-matching task", "Measures visual discrimination using a lower-production-demand activity before independent writing.", "First-attempt accuracy, completion time and repeated errors."),
        ("Finger tracing", "Creates an active writing interaction and the authentic application input required by the CNN.", "Classifier agreement, confidence, child-trace accuracy and usability. Geometric guides are approximate fallbacks."),
        ("Grade 2 pillam instruction", "Extends basic letter recognition to the vowel-sign knowledge required for Sinhala word construction.", "Educator sequencing review and pillam-identification accuracy."),
        ("Pillam fill-in", "Tests contextual application by selecting the missing pillam within a word rather than simple recall.", "Per-item completion accuracy and confusion analysis."),
        ("Pillam naming", "Measures connection between a displayed glyph and its linguistic name.", "Per-pillam accuracy, repeated mistakes and response time."),
        ("Child-centred visuals", "Large text, colour, icons and touch targets are design hypotheses intended to reduce reading and motor demands.", "Tap-error rate, readability, task completion and observation. Suitability must be tested rather than assumed."),
        ("Stars, streaks and confetti", "Provide immediate motivational feedback and encourage repeated practice without replacing educational feedback.", "Return rate, sessions completed and distraction reports. Achievements are not implemented and reading points remain zero."),
    ], font_size=8.3)

    doc.add_heading("6. Reading, Comprehension and Adaptation Matrix", level=1)
    add_table(doc, ["Feature", "Research justification and role", "Evaluation and honest boundary"], [
        ("Grade-adapted story reader", "Applies grade-level decisions within an actual reading workflow using shorter/larger Grade 1 pages and longer Grade 2 passages.", "Readability review, completion and help requests. The current library contains only three prototype stories."),
        ("Text-to-speech", "Provides an audio model and decoding support when the learner cannot independently read a word or passage.", "Usage, task completion and educator audio-quality ratings. It does not prove independent reading."),
        ("Speech-to-text/WPM", "Explores fluency measurement by dividing recognised word count by elapsed time.", "Compare with human transcript and timing. It is approximate WPM, not pronunciation or reading accuracy, because expected text is not compared."),
        ("Vocabulary helper", "Addresses the difference between decoding a word and understanding it through bilingual meanings and examples.", "Lookup use, subsequent recall and educator review. The curated database is small, unknown words use a generic fallback, and story word-tap content is placeholder."),
        ("Comprehension quizzes", "Measures whether a learner understood a story and provides scores by question type for progress and adaptation.", "Item accuracy, reliability and educator validation. Generation is rule/template based; open responses currently count as correct and must not be used as objective scores."),
        ("Difficulty-matched recommendations", "Uses the text classifier to select stories whose estimated difficulty matches the learner's Grade 1/2 level.", "Teacher appropriateness, subsequent quiz results and comparison with authored tags. Limited by the small classifier corpus and story pool."),
        ("Performance recommendation", "Turns comprehension results into an actionable easier/same/challenge decision instead of displaying a score only.", "Acceptance and subsequent performance. Current rules are mainly comprehension thresholds; this is not a trained recommender."),
        ("Continue reading", "Intended to preserve continuity for partially completed stories.", "Resume accuracy and completion. Mid-story autosave is not implemented, so the section currently has no reliable in-progress evidence."),
    ], font_size=8.3)

    doc.add_heading("7. Evidence, Stakeholder and Platform Feature Matrix", level=1)
    add_table(doc, ["Feature", "Research justification and role", "Evaluation and honest boundary"], [
        ("Progress tracking", "Transforms isolated task results into longitudinal evidence across practice, reading sessions and quizzes.", "Data accuracy, completeness and trend validity. Vocabulary retention, pronunciation accuracy and achievements are currently zero/empty."),
        ("Weekly weak-area insight", "Uses transparent thresholds to identify the weakest practiced item or task and recommend focused practice.", "Agreement with teacher judgement and false-warning rate. It is rule based, not machine learning."),
        ("Student profiles", "Separates records for different children and controls grade, language and daily-goal settings.", "Data-isolation tests, profile-selection usability and goal completion. Goals are design interventions, not proven habit improvement."),
        ("Parent/teacher roles", "Recognises that young learners require adult supervision and different stakeholders require different access.", "Role-routing and authorisation tests. Adult monitoring does not replace teacher judgement."),
        ("Teacher join code", "Provides limited sharing of one student's record without making the student collection publicly browsable.", "Valid/invalid code, unauthorised access and data-isolation tests. Detailed teacher analytics remain incomplete."),
        ("Firebase persistence", "Stores longitudinal sessions and account ownership so research evidence survives app restarts and can be audited.", "Write/read consistency, security-rule tests and failure handling. Android is configured; web/iOS options remain placeholders and connectivity is required."),
        ("Adult generative-AI assistant", "Translates stored progress into accessible explanations and suggests existing activities. It is a support interface, not a core evaluated model.", "Factuality, hallucination, refusal, language and usefulness. The exposed client API credential must be rotated and moved behind a protected service."),
        ("Chat history", "Allows an adult to revisit explanations and maintain continuity across discussions.", "Retrieval and access-control tests. Retention and privacy are required because conversations may contain child-related information."),
        ("Five-tab navigation", "Separates Home, Learn, AI Help, Progress and Profile to reduce route complexity for children and adults.", "Task-finding time, navigation errors and usability ratings."),
        ("Flutter/TFLite/Firebase architecture", "Balances cross-platform delivery, local inference and authorised cloud persistence in one testable architecture.", "Build success, device compatibility, latency, offline behaviour, security and failure recovery."),
    ], font_size=8.2)


def evidence_control(doc):
    doc.add_heading("8. Evaluation Plan by Feature Type", level=1)
    add_table(doc, ["Feature type", "Primary measures", "Evidence source", "Claim permitted"], [
        ("CNN model", "Accuracy, macro-F1, recall, confusion, top-k and calibration.", "Frozen validation/test data and reproducible notebook.", "Reference-data classification performance."),
        ("CNN integration", "Mapping correctness, confidence rule, latency and failure behaviour.", "Automated tests and Android device logs.", "Technical mobile feasibility."),
        ("Text classifier", "Balanced accuracy, macro-F1, baseline improvement and educator agreement.", "Expanded labelled Grade 1/2 corpus.", "Grade discrimination within evaluated data."),
        ("Learning tasks", "Accuracy, attempts, completion time and error patterns.", "Persisted task attempts.", "Application task performance, not school achievement."),
        ("Reading/quiz", "Completion, correct WPM, quiz reliability and question-type results.", "Sessions, quiz attempts and human checks.", "Observed application reading/comprehension performance."),
        ("Recommendations", "Teacher appropriateness, acceptance and subsequent performance.", "System logs and educator review.", "Rule/model recommendation usefulness within evaluated stories."),
        ("Usability", "Task success, time, errors and age-appropriate satisfaction.", "Ethics-approved moderated sessions.", "Usability for the sampled participants."),
        ("Educational effectiveness", "Pre/post learning outcomes and appropriate statistical analysis.", "Ethics-approved controlled field study.", "Only after the future study; not supported currently."),
    ], font_size=8.5)

    doc.add_heading("9. Claims to Avoid and Defensible Replacements", level=1)
    add_table(doc, ["Avoid this statement", "Use this defensible statement"], [
        ("The CNN has 94% test accuracy.", "The CNN achieved approximately 94% validation accuracy on reference data."),
        ("The application accurately recognises children's handwriting.", "Child-trace field accuracy has not yet been established."),
        ("All 454 letters work in the app.", "The model has 454 outputs, while the current UI exposes five letters and requires Unicode mapping validation."),
        ("The system measures pronunciation accuracy.", "The current speech feature measures approximate recognised-word rate and WPM."),
        ("The text classifier is production ready.", "It is an interpretable proof of concept fitted on 26 samples."),
        ("The quiz is generated by advanced AI.", "The quiz uses grade-aware rules and templates; open answers are not objectively scored."),
        ("The recommender uses all learner metrics.", "The current recommendation rule is mainly driven by comprehension thresholds."),
        ("The application proves literacy improvement.", "It is a developed research artifact; effectiveness requires an ethics-approved outcome study."),
        ("The Gemini assistant is the research model.", "Gemini is an adult-support feature; the CNN and text classifier are the research models."),
        ("All features are complete.", "The prototype is implemented with documented integration, security, data and evaluation gaps."),
    ], font_size=8.5)

    doc.add_heading("10. Priority Actions before Supervisor Demonstration", level=1)
    add_numbered(doc, [
        "Complete and test the numeric-class-ID-to-Sinhala-Unicode mapping for all five exposed tracing letters.",
        "Recover and archive the exact notebook, dataset version, seed and environment that produced the shipped 64×64 TFLite model.",
        "Rotate the exposed generative-AI API credential and move calls behind a protected backend or proxy with restrictions and usage limits.",
        "Expand and independently label the Grade 1/2 text corpus; evaluate it against a majority baseline using stratified procedures.",
        "Remove short-response items from objective quiz scores until a rubric or validated scoring method exists.",
        "Replace placeholder story word-help content and persist vocabulary learning if vocabulary metrics are displayed.",
        "Compare recognised speech with expected text before reporting reading accuracy or pronunciation results.",
        "Connect or remove the pre-reading and continue-reading features so every demonstrated feature has a reachable, functioning workflow.",
        "Correct static-analysis findings and replace the obsolete default widget test with model, mapping, task and critical navigation tests.",
        "Obtain supervisor and ethics approval before collecting images, audio or learning outcomes from children.",
    ])


def viva_and_signoff(doc):
    doc.add_heading("11. Supervisor/Viva Answer Scripts", level=1)
    add_callout(doc, "Thirty-second system structure", "The project has three layers. The core research layer contains the on-device Sinhala letter CNN and Grade 1/2 text-difficulty classifier. The educational-delivery layer contains letters, tracing, pillam, stories, audio, vocabulary and quizzes. The evidence-and-stakeholder layer contains progress tracking, recommendations, parent/teacher access and adult AI explanations. The supporting features create a realistic workflow for demonstrating and evaluating the two research models.")
    add_callout(doc, "Why so many supporting features?", "The supporting features are not claimed as separate innovations. They provide authentic model inputs, turn model outputs into educational actions, record measurable outcomes, and allow responsible adults to review evidence. Without this layer, the work would be two isolated model demonstrations rather than an evaluated mobile research artifact.")
    add_callout(doc, "How do you know a feature is necessary?", "A feature remains in scope only when it directly implements a core model, supplies its input, operationalises its output, collects evaluation evidence, supports the Grade 1–2 learner, or enables authorised adult interpretation. Features without one of these roles should be removed or identified as future work.")
    add_callout(doc, "What is the strongest honest contribution?", "The strongest current contribution is the integration of an on-device Sinhala character classifier with a Grade 1/2 bilingual learning workflow. The text-difficulty classifier is a second, interpretable low-resource proof of concept. The work is strengthened—not weakened—by separating reference validation, application integration and future child-field evidence.")

    doc.add_heading("12. Supervisor Decision Record", level=1)
    add_body(doc, "Complete this table only after the supervisor gives actual feedback. Do not invent approval or decisions.")
    add_table(doc, ["Discussion point", "Supervisor decision / feedback", "Student action and deadline"], [
        ("Confirm two core research components", "", ""),
        ("Confirm supporting-feature categorisation", "", ""),
        ("Confirm Grade 1/2 scope", "", ""),
        ("Confirm feature evaluation measures", "", ""),
        ("Confirm treatment of current limitations", "", ""),
        ("Confirm priority fixes before Results/Discussion", "", ""),
    ])
    add_body(doc, "Supervisor name: ______________________________   Signature: ____________________   Date: ____________")
    add_body(doc, "Student signature: ______________________________   Date: ____________________")


def references(doc):
    doc.add_heading("References — IEEE Style", level=1)
    refs = [
        '[1] UNICEF Sri Lanka, “MOE and UNICEF spearhead national initiative to recover lost learning for 1.6 million primary school children across Sri Lanka,” Aug. 16, 2023. [Online]. Available: https://www.unicef.org/srilanka/press-releases/moe-and-unicef-spearhead-national-initiative-recover-lost-learning-16-million.',
        '[2] G. F. Bautista, P. Ghesquière, and J. Torbeyns, “Stimulating preschoolers’ early literacy development using educational technology: A systematic literature review,” Int. J. Child-Comput. Interact., vol. 39, Art. no. 100620, 2024, doi: 10.1016/j.ijcci.2023.100620.',
        '[3] P. B. Gough and W. E. Tunmer, “Decoding, reading, and reading disability,” Remedial Spec. Educ., vol. 7, no. 1, pp. 6–10, 1986, doi: 10.1177/074193258600700104.',
        '[4] J. Mariyathas, V. Shanmuganathan, and B. Kuhaneswaran, “Sinhala handwritten character recognition using convolutional neural network,” in Proc. ICITR, 2020, pp. 1–6, doi: 10.1109/ICITR51448.2020.9310914.',
        '[5] W. V. S. K. Wasalthilake and T. Thangathurai, “Improved handwritten character recognition for Sinhala language based on convolutional neural networks,” in Proc. IEEE I2CT, 2022, doi: 10.1109/I2CT54291.2022.9824233.',
        '[6] M. L. Karunarathne, C. P. Wijesiriwardana, K. M. I. Nishantha, and W. G. C. W. Kumara, “Efficiency and accuracy in Sinhala handwritten character recognition: A Gabor-initialized CNN perspective,” Sri Lankan J. Technol., vol. 5, no. 1, pp. 13–24, 2024.',
        '[7] J. M. Imperial and E. Kochmar, “Automatic readability assessment for closely related languages,” in Findings of ACL, 2023, pp. 5371–5386, doi: 10.18653/v1/2023.findings-acl.331.',
        '[8] T. Naous, M. J. Ryan, A. Lavrouk, M. Chandra, and W. Xu, “ReadMe++: Benchmarking multilingual language models for multi-domain readability assessment,” in Proc. EMNLP, 2024, pp. 12230–12266, doi: 10.18653/v1/2024.emnlp-main.682.',
        '[9] F. Liu, T. Jin, and J. S. Y. Lee, “Automatic readability assessment for sentences: Neural, hybrid and large language models,” Lang. Resources Eval., vol. 59, pp. 2265–2296, 2025, doi: 10.1007/s10579-024-09800-5.',
        '[10] J. Lee et al., “On-device neural net inference with mobile GPUs,” in Workshop Efficient Deep Learning for Computer Vision, CVPR, 2019. [Online]. Available: https://arxiv.org/abs/1907.01989.',
        '[11] C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger, “On calibration of modern neural networks,” in Proc. ICML, vol. 70, 2017, pp. 1321–1330.',
        '[12] K. Peffers, T. Tuunanen, M. A. Rothenberger, and S. Chatterjee, “A design science research methodology for information systems research,” J. Manage. Inf. Syst., vol. 24, no. 3, pp. 45–77, 2007, doi: 10.2753/MIS0742-1222240302.',
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Pt(20)
        p.paragraph_format.first_line_indent = Pt(-20)
        p.paragraph_format.space_after = Pt(5)


def build():
    doc = setup_document()
    doc.sections[0].header.paragraphs[0].text = "ReadBuddy AI — Supervisor Feature Justification"
    doc.sections[0].header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title = "ReadBuddy AI — Research Justification for Every System Feature"
    doc.core_properties.author = "L. M. I. Silva"
    doc.core_properties.subject = "Supervisor Discussion and Viva Preparation — Student ID 28277"
    title_page(doc)
    front_matter(doc)
    research_structure(doc)
    core_features(doc)
    feature_matrices(doc)
    evidence_control(doc)
    viva_and_signoff(doc)
    references(doc)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
