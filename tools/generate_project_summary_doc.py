from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
PATH = OUT / "ReadBuddy_AI_Project_Summary_Aim_Target_Coverage.docx"
NAVY = "17365D"
BLUE = "2E75B6"
TEAL = "2A9D8F"
CORAL = "E76F51"
LIGHT = "F5F7FA"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.add_run(item)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        shade(t.rows[0].cells[i], NAVY)
        p = t.rows[0].cells[i].paragraphs[0]
        r = p.add_run(header)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        r.font.name = "Times New Roman"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, value in enumerate(row):
            if ri % 2:
                shade(cells[ci], LIGHT)
            cells[ci].text = str(value)
            cells[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[ci].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9.5)
    doc.add_paragraph()
    return t


def callout(doc, title, text):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    shade(c, "DFF2EF")
    p = c.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(text)
    doc.add_paragraph()


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21)
    sec.top_margin = Cm(2.3)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans Sinhala")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in [("Title", 22, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 13, BLUE)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    header.text = "ReadBuddy AI — Project Summary"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(107, 114, 128)
    add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("READBUDDY AI")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Project Summary, Aim, Target Users and Overall System Coverage")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    for text in [
        "AI-Enhanced Bilingual Reading and Comprehension Assistant",
        "Grade 1 and Grade 2 Primary School Students",
        "Mahima Induvara",
        "BSc (Hons) in Software Engineering — Final Year Research Project",
        "2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(16 if text == "Mahima Induvara" else 8)
        rr = p.add_run(text)
        rr.font.name = "Times New Roman"
        rr.font.size = Pt(13 if text == "Mahima Induvara" else 11)
        rr.bold = text == "Mahima Induvara"
    doc.add_page_break()

    doc.add_heading("1. Project Overview", level=1)
    body(doc, "ReadBuddy AI is a bilingual Sinhala and English mobile application designed to support reading, writing and comprehension development among Grade 1 and Grade 2 primary school students in Sri Lanka. The application combines curriculum-focused learning activities with artificial intelligence and provides letter learning, Sinhala letter tracing, word-building activities, reading exercises, text-to-speech support, comprehension quizzes, vocabulary assistance, progress tracking and reading recommendations.")
    body(doc, "The principal research contribution is the design and integration of two original machine-learning components: an on-device convolutional neural network for Sinhala handwritten-letter recognition and a lightweight classifier that predicts whether Sinhala reading content is more appropriate for Grade 1 or Grade 2. Flutter and Dart provide the mobile interface, TensorFlow Lite executes letter recognition locally, and Firebase supports authentication and authorized progress storage.")

    doc.add_heading("2. Project Aim", level=1)
    callout(doc, "Main aim", "To design, develop and evaluate an AI-enhanced bilingual mobile reading and comprehension assistant that supports Grade 1 and Grade 2 students through interactive literacy activities, on-device Sinhala letter recognition and automated Sinhala text-difficulty classification.")

    doc.add_heading("3. Target Users", level=1)
    doc.add_heading("3.1 Primary Users", level=2)
    bullets(doc, [
        "Grade 1 students learning Sinhala letters, simple words, letter tracing and basic reading skills.",
        "Grade 2 students developing pillam knowledge, word formation, sentence reading, vocabulary and comprehension skills.",
    ])
    doc.add_heading("3.2 Secondary Users", level=2)
    bullets(doc, [
        "Parents who want to review their child's learning activities and progress.",
        "Teachers who need additional digital support for early-literacy instruction.",
        "Researchers interested in Sinhala educational technology and low-resource-language machine learning.",
    ])
    body(doc, "ReadBuddy AI is intended to support students, parents and teachers. It is not designed to replace classroom instruction or provide medical, psychological or learning-disability diagnoses.")

    doc.add_heading("4. Problem Addressed", level=1)
    body(doc, "Many existing early-literacy applications primarily support English and other high-resource languages. Sinhala-speaking learners have fewer intelligent learning tools capable of providing automated, locally relevant feedback. ReadBuddy AI responds to two specific problems:")
    bullets(doc, [
        "Ordinary tracing applications can record a child's drawing but may not determine whether the expected Sinhala letter was produced correctly.",
        "Reading passages are commonly assigned to a grade manually, without an automated mechanism for checking whether their measurable linguistic difficulty is consistent with Grade 1 or Grade 2.",
    ])

    doc.add_heading("5. Overall System Coverage", level=1)
    table(doc, ["Area", "Coverage"], [
        ("Account and profile management", "Parent authentication, student profile setup, grade selection, language selection and progress persistence."),
        ("Grade 1 learning", "Letter introduction, pronunciation, tracing, recognition, matching, simple vocabulary, short reading and basic comprehension."),
        ("Grade 2 learning", "Pillam learning, identification and fill-in activities, word construction, longer passages, vocabulary and literal/basic inferential comprehension."),
        ("Reading support", "Sinhala/English content, text-to-speech, word assistance, comprehension quizzes, adaptive settings and recommendations."),
        ("Progress monitoring", "Learning attempts, quiz outcomes, reading sessions, model confidence and activity history."),
        ("Adult support", "Parent/teacher progress views, dashboard functions and an adult-facing AI Assistant."),
    ])

    doc.add_heading("6. Grade 1 Module Coverage", level=1)
    bullets(doc, [
        "Sinhala letter introduction with names, examples and audio pronunciation.",
        "Interactive tracing canvas with visual guidance.",
        "On-device classification of the captured trace.",
        "Letter-recognition and letter-matching activities.",
        "Simple Grade 1 words, short reading content and basic comprehension activities.",
    ])
    callout(doc, "Current prototype boundary", "The current AI tracing screen exposes five Sinhala letters: අ, ආ, ක, ග and ස. It should therefore be described as a working Grade 1 research prototype rather than complete alphabet or curriculum coverage.")

    doc.add_heading("7. Grade 2 Module Coverage", level=1)
    bullets(doc, [
        "Introduction to Sinhala pillam and word-building patterns.",
        "Pillam identification, naming and missing-pillam activities.",
        "Longer reading passages and expanded vocabulary.",
        "Literal comprehension questions and basic why/predict questions.",
        "Grade-appropriate reading recommendations and progress tracking.",
    ])

    doc.add_heading("8. Research Component 1 — Sinhala Letter Recognition", level=1)
    body(doc, "The first component is a custom convolutional neural network trained to classify isolated handwritten Sinhala characters. The model accepts 64×64 grayscale images, produces a prediction across 454 character/modification classes and is deployed through TensorFlow Lite. Inference runs directly on the mobile device, avoiding a cloud request for every trace.")
    table(doc, ["Evaluation metric", "Reported value"], [
        ("Validation accuracy", "94.04%"),
        ("Macro precision", "0.944"),
        ("Macro recall", "0.940"),
        ("Macro F1-score", "0.940"),
        ("Top-2 accuracy", "96.35%"),
        ("Top-3 accuracy", "97.12%"),
        ("Expected Calibration Error", "0.084"),
    ])
    body(doc, "These figures demonstrate strong performance on the reference validation dataset but do not establish equal performance on real children's traces. A current integration gap is that the deployed label file contains numeric class IDs from 1 to 454, while the user interface expects Sinhala Unicode letters. The mapping must be completed and tested before reliable end-to-end letter correctness can be claimed.")

    doc.add_heading("9. Research Component 2 — Text-Difficulty Classification", level=1)
    body(doc, "The second component is a binary logistic-regression classifier that predicts whether Sinhala text is closer to Grade 1 or Grade 2 difficulty. It uses word count, average word length and average sentence length. The model was trained using 26 application-content samples: 9 Grade 1 and 17 Grade 2 samples.")
    table(doc, ["Characteristic", "Details"], [
        ("Model", "Binary logistic regression"),
        ("Classes", "Grade 1 and Grade 2"),
        ("Features", "Word count, average word length and average sentence length"),
        ("Training samples", "26 total: 9 Grade 1 and 17 Grade 2"),
        ("Reported performance", "69.2% training accuracy"),
        ("Deployment", "Hand-written Dart sigmoid calculation; no mobile ML runtime required"),
    ])
    body(doc, "Because no independent test set was used and the dataset is very small, 69.2% must be described as training accuracy. This component is an exploratory proof-of-concept that validates the classification and integration pipeline. A larger, balanced and independently teacher-labelled Sinhala corpus is required before making a generalization claim.")

    doc.add_heading("10. Progress and Adult-Support Coverage", level=1)
    bullets(doc, [
        "Storage of completed activities, quiz results and reading sessions.",
        "Recording of correct/incorrect attempts and available model confidence information.",
        "Progress summaries and recommended activities for the learner's level.",
        "Authorized parent/teacher access to relevant student information.",
        "An adult-facing AI Assistant for educational guidance and application support.",
    ])
    body(doc, "The AI Assistant is intended for parents and teachers rather than unsupervised use by children. Progress records show application activity; they must not be presented as evidence of improved academic achievement without a formal learning-outcome study.")

    doc.add_heading("11. Technology Coverage", level=1)
    table(doc, ["Technology", "Purpose"], [
        ("Flutter and Dart", "Mobile interface, learning workflows and local text-model inference."),
        ("TensorFlow/Keras", "Training and evaluating the Sinhala CNN."),
        ("TensorFlow Lite / tflite_flutter", "On-device offline letter recognition."),
        ("Scikit-learn", "Training the logistic-regression text classifier."),
        ("Firebase Authentication", "Account and identity management."),
        ("Cloud Firestore", "Student profiles, progress and conversation persistence."),
        ("Cloud Functions and Gemini", "Adult-facing AI Assistant services."),
        ("Text-to-speech and speech-to-text", "Reading support and voice-enabled interactions."),
        ("Android", "Primary tested mobile deployment target."),
    ])

    doc.add_heading("12. Research Objectives", level=1)
    numbered = [
        "To identify gaps in Sinhala early-literacy applications, handwriting recognition and text-difficulty assessment.",
        "To analyze suitable machine-learning methods and mobile deployment approaches for the identified problems.",
        "To design and develop an on-device Sinhala letter-recognition CNN and a Grade 1/2 text-difficulty classifier.",
        "To evaluate the components and their application integration using transparent quantitative and functional evidence.",
    ]
    for item in numbered:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.add_run(item)

    doc.add_heading("13. Project Scope", level=1)
    table(doc, ["In scope", "Out of scope"], [
        ("Grade 1 and Grade 2 learners", "Grade 3 or higher-grade classification"),
        ("Sinhala and English interfaces", "A universal multilingual learning system"),
        ("Isolated Sinhala letter classification", "Full pages, worksheets, connected handwriting or OCR segmentation"),
        ("Five-letter tracing research prototype", "Complete Sinhala alphabet stroke-path coverage"),
        ("Binary Grade 1/2 text classification", "Clinical diagnosis, dyslexia detection or semantic comprehension diagnosis"),
        ("Progress and activity monitoring", "Proof of learning improvement without a controlled participant study"),
        ("Android mobile deployment", "Replacement of teachers or unsupervised autonomous teaching"),
    ])

    doc.add_heading("14. Main Strengths", level=1)
    bullets(doc, [
        "Addresses Sinhala, an under-resourced language in educational AI.",
        "Combines reading, writing, comprehension and progress support within one application.",
        "Uses original trained models rather than relying only on third-party AI APIs.",
        "Runs letter recognition on-device for offline access and reduced image transfer.",
        "Reports multiple CNN evaluation metrics rather than accuracy alone.",
        "Maintains a clear Grade 1–2 target and acknowledges evidence limitations honestly.",
    ])

    doc.add_heading("15. Current Limitations", level=1)
    bullets(doc, [
        "The CNN has been evaluated on reference dataset images rather than real children's traces.",
        "The numeric CNN classes are not yet fully mapped to Sinhala Unicode letters.",
        "Only five letters are exposed in the current tracing prototype.",
        "The tracing guides are approximate visual paths rather than verified letter-specific stroke-order data.",
        "The text classifier contains only 26 samples and has no held-out test set.",
        "No classroom or child-participant evaluation has been conducted.",
        "Activity progress does not by itself demonstrate educational improvement.",
    ])

    doc.add_heading("16. Overall Contribution", level=1)
    body(doc, "ReadBuddy AI contributes a practical research prototype combining mobile software engineering, computer vision and natural-language processing for Sinhala early literacy. Its strongest current contribution is the development and on-device integration of a high-performing Sinhala character classifier. Its second contribution is an interpretable Grade 1/2 Sinhala text-difficulty pipeline that demonstrates how reading content could eventually be checked and recommended automatically.")
    body(doc, "Overall, the project demonstrates how locally relevant artificial intelligence can be incorporated into a bilingual early-learning application while remaining transparent about small datasets, incomplete label mapping and the absence of classroom-based evaluation. It provides a foundation for future work involving complete curriculum coverage, teacher-labelled reading corpora and ethically approved evaluation with real learners.")

    doc.add_heading("17. Short Presentation Summary", level=1)
    callout(doc, "One-minute explanation", "ReadBuddy AI is a Sinhala and English mobile reading assistant developed for Grade 1 and Grade 2 students. It provides letter learning, tracing, word-building, reading, comprehension and progress-monitoring activities. The main research contribution is the integration of two original models: an on-device CNN for Sinhala handwritten-letter recognition and a lightweight classifier for distinguishing Grade 1 and Grade 2 Sinhala text. The system supports students while allowing parents and teachers to monitor learning activities. It is a research prototype, not a replacement for teachers or a diagnostic system. The CNN has strong validation performance, while the text classifier establishes an initial pipeline that requires a larger expert-labelled dataset for future improvement.")

    doc.core_properties.title = "ReadBuddy AI Project Summary — Aim, Target and Overall Coverage"
    doc.core_properties.author = "Mahima Induvara"
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(PATH)
    print(PATH)


if __name__ == "__main__":
    build()
