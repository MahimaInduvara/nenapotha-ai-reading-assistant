from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from generate_thesis_breakdown import (
    BLUE,
    CORAL,
    MID_GREY,
    NAVY,
    PALE_ORANGE,
    PALE_TEAL,
    add_body,
    add_bullets,
    add_callout,
    add_numbered,
    add_table,
    add_toc,
    setup_document,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
DOCX = OUT / "ReadBuddy_AI_Zero_Cost_Development_Roadmap_28277.docx"


def title_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(65)
    r = p.add_run("READBUDDY AI")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(27)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Zero-Cost Stage-by-Stage Development Roadmap")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("From Current Prototype to Final Thesis Submission")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(CORAL)

    for text, size, bold in [
        ("AI-Enhanced Bilingual Reading and Comprehension Assistant for Grade 1 and Grade 2 Students", 14, True),
        ("On-Device Sinhala Letter Recognition and Sinhala Text-Difficulty Classification", 13, False),
        ("L. M. I. Silva", 14, True),
        ("Student ID: 28277", 12, False),
        ("BSc (Hons) in Software Engineering — Final Year Research Project", 11, False),
        ("NSBM Green University — Faculty of Computing", 11, False),
        ("Prepared: 21 August 2026", 11, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(24 if text == "L. M. I. Silva" else 8)
        rr = p.add_run(text)
        rr.font.name = "Times New Roman"
        rr.font.size = Pt(size)
        rr.bold = bold

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    rr = p.add_run("Supervisor approval: ______________________________________________")
    rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor.from_string(MID_GREY)
    doc.add_page_break()


def front_matter(doc):
    doc.add_heading("How to Use This Roadmap", level=1)
    add_body(doc, "This roadmap gives the required order for completing ReadBuddy AI without purchasing software, hosting, compute or paid AI services. Each stage has a dependency, task list, zero-cost toolset, evidence output and completion gate. A later stage must not be treated as complete until the previous gate has passed. The schedule assumes approximately eight to ten weeks of part-time work; it may be compressed by reducing optional product work, but research-integrity stages should not be skipped.")
    add_callout(doc, "Main strategy", "Finish correctness, reproducibility and evaluation before adding more features. Full marks are more likely to come from a small, complete and defensible research contribution than from a large application containing unfinished or unsupported functions.", PALE_TEAL)
    add_callout(doc, "Zero-cost rule", "Use the existing laptop, Android emulator or an already available Android phone, Flutter/Dart, Python, scikit-learn, TensorFlow, TensorFlow Lite, Git, Google Colab free tier and Firebase Spark quotas. Do not upgrade to Firebase Blaze, purchase cloud compute, buy a generative-AI subscription, publish to Google Play, or purchase a domain merely for the thesis demonstration.", PALE_ORANGE)
    doc.add_heading("Table of Contents", level=1)
    add_toc(doc)
    doc.add_page_break()


def overview(doc):
    doc.add_heading("1. Master Development Map", level=1)
    add_table(doc, ["Order", "Stage", "Indicative time", "Main output / gate"], [
        ("0", "Freeze baseline and organise evidence", "1 day", "Recoverable project snapshot and issue register."),
        ("1", "Secure and clean the current application", "2–3 days", "No exposed secret; agreed Grade 1/2 scope; clean build baseline."),
        ("2", "Correct CNN-to-application integration", "3–5 days", "Five tracing letters pass end-to-end label tests."),
        ("3", "Reproduce and evaluate the CNN", "5–8 days", "Exact notebook, split, metrics, model card and device benchmark."),
        ("4", "Build the Grade 1/2 Sinhala text corpus", "7–14 days", "Balanced, independently labelled corpus and data card."),
        ("5", "Evaluate and select the text classifier", "4–7 days", "Baseline/model comparison, ablation and frozen final model."),
        ("6", "Complete essential learning workflows", "5–7 days", "No placeholder or falsely scored child-facing workflow."),
        ("7", "Complete progress and adult evidence views", "4–6 days", "Only real stored metrics shown; teacher review workflow works."),
        ("8", "Testing, performance and privacy verification", "5–7 days", "Passing automated tests, device logs and security checklist."),
        ("9", "Supervisor, expert and ethics-approved evaluation", "1–2 weeks", "Signed decisions, expert validation and approved user evidence."),
        ("10", "Write Results and Discussion", "5–7 days", "Every claim linked to generated evidence and limitations."),
        ("11", "Final packaging, rehearsal and submission", "3–4 days", "APK, source, evidence pack, thesis, slides and demo ready."),
    ], font_size=8.5)

    doc.add_heading("2. Zero-Cost Technology Plan", level=1)
    add_table(doc, ["Need", "Free solution", "Cost-control condition"], [
        ("Flutter application", "Flutter and Dart SDK; VS Code or Android Studio.", "Open-source development stack; no paid IDE plugin required [1]."),
        ("Android testing", "Android Emulator and an existing physical phone if available.", "Submit/install an APK directly; Google Play publication is unnecessary."),
        ("CNN training", "Local Python or Google Colab free runtime.", "Save checkpoints frequently because free Colab resources are limited and not guaranteed [2]."),
        ("Data analysis", "Python, pandas, NumPy, scikit-learn, SciPy, Matplotlib and Seaborn.", "All run locally or in a notebook."),
        ("Backend", "Firebase Authentication and Firestore on the Spark plan.", "Stay within quotas; no payment method or Blaze upgrade [3], [4]."),
        ("Generative assistant", "Existing local rule-based AIService; make Gemini optional or remove it from the final evaluated artifact.", "Do not depend on a paid API or client-embedded secret."),
        ("Version control", "Local Git; optional GitHub free repository with private-data exclusions.", "Never commit credentials, child data or consent records."),
        ("Forms and review", "University forms or privacy-reviewed Google Forms/Sheets using anonymous codes.", "Use only after supervisor/ethics approval; avoid identifiable child data."),
        ("Statistics", "Python/SciPy or JASP.", "Pre-register metrics; no commercial statistics package is needed."),
        ("Diagrams and reports", "draw.io, Matplotlib, LibreOffice or existing Microsoft Word.", "Export diagrams and tables into the thesis evidence pack."),
    ], font_size=8.5)


def add_stage(doc, number, title, purpose, tasks, outputs, gate, tools, avoid=None):
    doc.add_heading(f"Stage {number} — {title}", level=1)
    add_callout(doc, "Purpose", purpose, PALE_TEAL)
    doc.add_heading("Tasks in order", level=2)
    add_numbered(doc, tasks)
    doc.add_heading("Required outputs", level=2)
    add_bullets(doc, outputs)
    add_table(doc, ["Free tools", "Completion gate"], [(tools, gate)], font_size=9)
    if avoid:
        add_callout(doc, "Do not proceed if", avoid, PALE_ORANGE)


def stages_0_to_3(doc):
    add_stage(doc, "0", "Freeze the Baseline and Organise Evidence",
        "Create a recoverable starting point so later results, fixes and model versions can be compared rather than remembered informally.",
        [
            "Create a local Git repository if one does not already exist; add .gitignore entries for build output, API secrets, datasets containing personal data and generated temporary files.",
            "Create folders for docs, experiments, models, data_cards, model_cards, test_reports and demo_evidence.",
            "Record the current APK/build status, Flutter version, Dart version, dependency lock file, current model SHA-256 hash and class-label file hash.",
            "Capture the current flutter analyze and flutter test results without hiding failures.",
            "Create an issue register with severity, research impact, responsible action and status.",
            "Copy the final approved research aim, one research question and four objectives into a traceability file.",
        ],
        ["Baseline tag or zip", "Environment record", "Issue register", "Research-objective traceability sheet", "Initial analyzer/test reports"],
        "Every later artifact can be traced to a date/version, and the current state can be restored without deleting user work.",
        "Git, PowerShell, Flutter CLI, SHA-256 utilities and Markdown/Word.")

    add_stage(doc, "1", "Secure and Clean the Current Application",
        "Remove risks and scope contradictions before producing new research evidence.",
        [
            "Rotate the generative-AI key currently embedded in the Flutter client and remove all fallback secrets from source code.",
            "For the zero-cost final artifact, use the existing local rule-based AIService as the default adult assistant; treat external Gemini access as optional and outside the evaluated core.",
            "Confirm Grade 1 and Grade 2 only in constants, profile choices, content, labels, thesis and presentation.",
            "Remove or hide unsupported Grade 3, diagnosis, pronunciation-accuracy and proven-effectiveness claims.",
            "Resolve flutter analyze warnings and determine whether every informational finding will be fixed or documented.",
            "Replace the obsolete default counter widget test with a minimal app-start/navigation smoke test.",
            "Verify Android Firebase configuration and mark web/iOS as unsupported unless properly configured.",
        ],
        ["Secret-free source tree", "Scope-control checklist", "Clean or justified analyzer report", "Passing basic smoke test", "Updated risk register"],
        "No secret is embedded in the application, the build opens successfully, and every visible grade claim is Grade 1/2.",
        "Flutter CLI, Firebase Spark, local AIService and Git.",
        "A credential is still visible in source, the app cannot start, or Grade 3 remains in an active claim.")

    add_stage(doc, "2", "Correct CNN-to-Application Integration",
        "Make the model output, Unicode letters, UI verdict and saved progress agree end to end.",
        [
            "Recover the authoritative mapping from class IDs 1–454 to Sinhala Unicode classes; document its source and version.",
            "Replace numeric-only label comparison with a verified mapped Unicode comparison.",
            "Create unit tests for the five exposed tracing letters: අ, ආ, ක, ග and ස.",
            "Test high-confidence correct, high-confidence wrong, low-confidence, blank input, corrupted image, missing model and output-size mismatch cases.",
            "Ensure the on-screen verdict, score bar and persisted TaskAttempt always use the same final decision.",
            "Replace misleading fallback wording: when the CNN is unavailable, clearly state that automatic recognition is unavailable.",
            "Record a short screen video and test log for each of the five letters.",
        ],
        ["Verified mapping file", "Mapping/unit tests", "Five-letter integration matrix", "Failure-mode evidence", "Android demonstration video"],
        "All five exposed letters pass deterministic mapping and UI/storage agreement tests; no numeric ID is directly compared with a Unicode letter.",
        "Dart/Flutter tests, existing TFLite model, Android Emulator and screen recorder.",
        "Any letter can be visually marked correct while a different result is stored, or the label source is unverified.")

    add_stage(doc, "3", "Reproduce and Evaluate the CNN",
        "Turn the promising model into a reproducible research experiment with correctly named metrics and mobile evidence.",
        [
            "Recover or recreate the exact 64×64 custom-CNN training notebook that produced the shipped TFLite model.",
            "Create a deterministic train/validation/test split and record image counts per class; prevent duplicate or source-related leakage across splits.",
            "Train with fixed seeds and save complete logs, learning curves and model checkpoints.",
            "Report test accuracy, macro/weighted F1, per-class recall, confusion matrix, top-1/top-3 accuracy and confidence calibration.",
            "Compare the custom CNN against one feasible baseline such as MobileNetV2/MobileNetV3 or a smaller five-class model using identical splits.",
            "Measure TFLite size, load success, median latency and p95 latency on the emulator and an existing Android phone if available.",
            "Create a Model Card containing purpose, intended users, data, metrics, limitations, ethical risks and unsupported uses.",
        ],
        ["Executable notebook", "Frozen split manifest", "Results CSV", "Training/validation curves", "Confusion matrix", "Calibration plot", "Latency report", "Model Card", "Final model hash"],
        "A fresh notebook run can reproduce the reported architecture and metrics, and the exact evaluated model is the model shipped in the app.",
        "Google Colab free or local Python, TensorFlow, scikit-learn, Matplotlib and TensorFlow Lite.",
        "The notebook still describes a different architecture, the test set influenced model selection, or only overall accuracy is available.")


def stages_4_to_7(doc):
    add_stage(doc, "4", "Build the Grade 1/2 Sinhala Text Corpus",
        "Replace the 26-sample proof of concept with traceable, balanced and independently labelled evidence.",
        [
            "Write a one-page Grade 1/2 annotation guideline covering passage length, word familiarity, sentence structure, curriculum alignment and disagreement handling.",
            "Collect legally usable curriculum-aligned passages and record source, grade, topic and permission status.",
            "Remove duplicates and near-duplicates; create anonymous stable text IDs.",
            "Aim first for at least 100 balanced passages and continue toward 200–400 if feasible; use a learning curve rather than claiming that a fixed count guarantees adequacy.",
            "Ask at least two educators to label independently; use a third review or adjudication for disagreements.",
            "Calculate raw agreement and Cohen's kappa; retain original labels and adjudicated labels.",
            "Create a Data Card documenting source, balance, cleaning, exclusions, annotation process, limitations and permitted use.",
        ],
        ["Annotation guide", "Balanced corpus CSV/JSON", "Source register", "Agreement report", "Adjudication log", "Data Card"],
        "The corpus has traceable sources, no known duplicates across evaluation splits, two independent labels and documented agreement.",
        "LibreOffice/Google Sheets, Python, pandas, scikit-learn and university-approved educator contact.",
        "Only the developer labelled the data, Grade 1/2 are strongly imbalanced, or copied text has unclear permission.")

    add_stage(doc, "5", "Evaluate and Select the Text Classifier",
        "Select an interpretable final model through comparison rather than assumption.",
        [
            "Freeze the cleaned corpus before modelling and keep an untouched final test set if the corpus is large enough; otherwise use nested or repeated stratified cross-validation.",
            "Establish a majority-class and simple length-rule baseline.",
            "Compare logistic regression, linear SVM, random forest and a character n-gram linear model using identical folds.",
            "Add interpretable Sinhala-related features such as sentence count, long-word ratio, unique-word ratio, vowel-sign/pillam frequency and curriculum-word coverage where reliable.",
            "Run an ablation study: surface features only, Sinhala-aware features only and combined features.",
            "Report macro-F1, balanced accuracy, per-grade recall, confusion matrix, confidence intervals and inference time.",
            "Select the smallest defensible model, freeze weights/features and update the Flutter implementation plus tests.",
        ],
        ["Baseline table", "Cross-validation results", "Feature-ablation table", "Learning curve", "Error analysis", "Frozen final parameters", "Text-model Model Card", "Flutter parity tests"],
        "The selected model beats the baseline consistently, its Flutter predictions match the Python reference cases, and its limitations are stated.",
        "Python, pandas, scikit-learn, SciPy, Matplotlib, Colab free or local CPU.",
        "Only training accuracy is reported or a complex model is selected without outperforming a simple baseline.")

    add_stage(doc, "6", "Complete Essential Learning Workflows",
        "Remove misleading or incomplete child-facing behaviour so the evaluated artifact matches the thesis description.",
        [
            "Connect the pre-reading sequence to navigation or remove it from the demonstrated scope.",
            "Make the reading completion button open the real QuizScreen rather than a placeholder message.",
            "Exclude short-answer questions from automatic numerical scoring until teacher/rubric assessment exists.",
            "Replace placeholder word-help images/meanings with the curated vocabulary helper or a clear not-available state.",
            "Implement mid-story autosave before showing Continue Reading, or hide the section.",
            "For speech, align recognised text with expected text before reporting correct WPM; otherwise label the value approximate and remove pronunciation-accuracy claims.",
            "Ensure Grade 1 and Grade 2 content, font size, question type and help behaviour follow the documented settings.",
        ],
        ["Reachability map", "Completed reading-to-quiz flow", "Correct scoring rules", "Real vocabulary help", "Autosave/resume tests", "Speech-measurement decision record"],
        "Every visible button has a real destination, every displayed score has an implemented calculation, and no placeholder is presented as a completed feature.",
        "Flutter/Dart, local curated content, SharedPreferences/Firestore within Spark quotas.",
        "A child can receive a false correct mark, placeholder definition or non-functional action.")

    add_stage(doc, "7", "Complete Progress and Adult Evidence Views",
        "Show only measurements supported by stored evidence and introduce transparent educator-in-the-loop review.",
        [
            "Display real practice accuracy, reading sessions, quiz performance, comprehension by type and correct/approximate WPM with clear labels.",
            "Keep vocabulary retention, pronunciation accuracy and achievements hidden until their data pipelines are implemented.",
            "Add minimum-evidence rules before producing weak-area insights and show 'insufficient evidence' for sparse data.",
            "Finish the teacher dashboard with linked-student activity, weak letters/pillam, comprehension trends and recent sessions.",
            "Create a teacher content-validation page: paste text, view predicted grade and features, approve/correct the grade and export anonymous corrections.",
            "Add an anonymous CSV export for research analysis; exclude names, email addresses and conversation content.",
            "Keep generative AI optional; default to local rule-based adult explanations to maintain zero cost and predictable evidence.",
        ],
        ["Evidence-to-dashboard mapping", "Insufficient-data states", "Teacher review workflow", "Anonymous export", "Role/access test matrix"],
        "Every displayed metric can be traced to stored source events, and the teacher can inspect/correct rather than blindly accept AI output.",
        "Flutter, Firestore Spark, local AIService, CSV export and Firebase security rules.",
        "A dashboard shows invented zeros as performance, labels a child at risk without validation, or exposes another student's data.")


def stages_8_to_11(doc):
    add_stage(doc, "8", "Testing, Performance and Privacy Verification",
        "Produce examiner-ready technical evidence that the application is correct, stable, efficient and appropriately protected.",
        [
            "Write unit tests for label mapping, preprocessing, text features, sigmoid output, recommendation thresholds and progress aggregation.",
            "Write widget tests for login/profile flow, Grade 1/2 learning navigation, tracing verdict, quiz scoring and insufficient-data states.",
            "Write integration scenarios covering model unavailable, offline mode, failed Firestore writes, restart/recovery and teacher authorisation.",
            "Test Firestore rules with owner, linked teacher, unrelated user and unauthenticated identities.",
            "Profile model load time, median/p95 inference, app memory and critical-screen responsiveness.",
            "Create a privacy inventory: each collected field, purpose, storage location, access, retention and deletion procedure.",
            "Run flutter analyze and flutter test; archive complete outputs and resolve all critical failures.",
        ],
        ["Unit/widget/integration report", "Firestore-rule report", "Performance table", "Privacy inventory", "Failure-recovery evidence", "Final analyzer report"],
        "All critical tests pass, no credential is exposed, access-control cases behave correctly, and latency is reported from repeatable runs.",
        "Flutter test, Firebase Emulator Suite, Android Profiler, PowerShell and local Python.",
        "The obsolete counter test still represents the test suite, Firestore rules are untested, or a critical failure is hidden.")

    add_stage(doc, "9", "Supervisor, Expert and Ethics-Approved Evaluation",
        "Validate content and usability without making unsupported claims or collecting child data without approval.",
        [
            "Discuss and obtain supervisor decisions on scope, model terminology, participant plan, metrics and storage of drawings/audio.",
            "First conduct an expert review with educators for content, grade labels, feedback wording and recommendation suitability.",
            "Apply required revisions and record a before/after change log.",
            "Submit ethics documents, information sheets, consent/assent materials and data-retention plan before involving children.",
            "After approval, conduct the sample and procedure agreed with the supervisor or power analysis; do not invent participant counts.",
            "Measure task completion, navigation errors, help requests, satisfaction and teacher/parent observations.",
            "If learning outcomes are assessed, use a justified pre/post or comparison design and report uncertainty and limitations.",
        ],
        ["Signed supervisor record", "Educator review forms", "Content-agreement results", "Ethics approval", "Consent/assent pack", "Anonymised usability dataset", "Change log"],
        "All human data were collected under approval and consent, participant records are anonymised, and the claimed result matches the actual design.",
        "University ethics process, digital forms, existing devices, LibreOffice/Sheets, Python/JASP.",
        "Ethics approval is pending, raw child data are identifiable, or a two-teacher preliminary survey is presented as population evidence.")

    add_stage(doc, "10", "Write Results and Discussion",
        "Convert generated technical and human evidence into transparent answers to the research question.",
        [
            "Freeze all final models, application version, datasets and experiment outputs before writing results.",
            "Present dataset composition and participant/expert characteristics without identifying individuals.",
            "Report CNN metrics, confusion, calibration, model comparison and mobile latency.",
            "Report text-classifier baselines, cross-validation, ablation, agreement and error examples.",
            "Report software-test outcomes and usability/expert results separately from model accuracy.",
            "Discuss why errors occurred, how findings compare with literature, threats to validity, ethical implications and generalisation limits.",
            "Answer each research-objective success indicator and the primary research question using explicit evidence references.",
        ],
        ["Results tables and figures", "Statistical outputs", "Objective-to-evidence matrix", "Threats-to-validity table", "Discussion comparison matrix", "Chapter drafts"],
        "Every numerical claim can be reproduced from an archived file and every conclusion is limited to the evaluated data/population.",
        "Word/LibreOffice, Python/Matplotlib, JASP and verified IEEE sources.",
        "A number cannot be traced to an output file, validation is called testing, or technical performance is presented as educational effectiveness.")

    add_stage(doc, "11", "Final Packaging, Rehearsal and Submission",
        "Deliver a reproducible thesis artifact and a demonstration that remains functional without paid services.",
        [
            "Build a signed or debug APK for direct installation; verify it on the emulator and an available Android phone.",
            "Create a demo account/data set containing no real child information and prepare an offline demonstration path.",
            "Package source, README, environment file template, notebooks, model/data cards, hashes, test reports and anonymised results.",
            "Update the thesis table of contents, lists, IEEE references, captions, page numbers and appendices.",
            "Prepare a 7–10 minute presentation: problem, gap, two models, architecture, methodology, results, limitations and contribution.",
            "Rehearse failure questions: small dataset, validation versus test, mapping, child privacy, AI assistant, generalisation and cost.",
            "Create two backups on separate existing storage locations and open every final file before submission.",
        ],
        ["APK", "Source archive", "Reproduction README", "Evidence pack", "Final thesis PDF/DOCX", "Slides", "Demo video", "Viva question sheet", "Two verified backups"],
        "The application can be demonstrated with no payment, no secret, no internet-dependent core model and no real child data; all submitted files open correctly.",
        "Flutter build, Android Emulator/device, screen recorder, Word/LibreOffice, local Git and existing storage.",
        "The demo depends on an external paid API, only one copy of evidence exists, or results cannot be matched to the submitted model.")


def schedule_and_control(doc):
    doc.add_heading("3. Suggested Ten-Week Execution Schedule", level=1)
    add_table(doc, ["Week", "Primary work", "End-of-week evidence"], [
        ("1", "Stages 0–1: baseline, security, scope and clean build.", "Versioned baseline, secret-free source, issue register and smoke test."),
        ("2", "Stage 2: Unicode mapping and five-letter CNN integration.", "Passing end-to-end tests and demo video."),
        ("3", "Stage 3: reproduce CNN and freeze data split.", "Executable notebook, split manifest and training outputs."),
        ("4", "Finish CNN evaluation and begin corpus collection.", "CNN results/model card and annotation guideline."),
        ("5", "Stage 4: educator annotation and agreement.", "Balanced corpus, kappa and data card."),
        ("6", "Stage 5: baselines, model comparison and Flutter parity.", "Final text model, ablation and tests."),
        ("7", "Stages 6–7: complete learning/progress workflows.", "No placeholders; teacher/evidence workflows operational."),
        ("8", "Stage 8: automated, security and performance testing.", "Test pack, privacy inventory and latency report."),
        ("9", "Stage 9 or supervisor-approved expert/user evaluation.", "Signed decision, approved anonymised evidence and change log."),
        ("10", "Stages 10–11: write, package and rehearse.", "Final chapters, evidence pack, APK, slides and backups."),
    ], font_size=8.5)

    doc.add_heading("4. Cost-Control Checklist", level=1)
    add_table(doc, ["Potential expense", "Zero-cost decision"], [
        ("Firebase Blaze/Cloud Functions", "Do not upgrade. Use Spark Auth/Firestore and local rules; external AI is optional."),
        ("Paid Gemini/OpenAI API", "Use the local rule-based assistant or disable adult generative chat in the evaluated build."),
        ("Paid GPU", "Use Colab free when available or train locally; save checkpoints and schedule shorter experiments."),
        ("Google Play developer fee", "Distribute the APK directly for supervisor/examiner testing."),
        ("Paid hosting/domain", "Not required; demonstrate locally or with direct APK/video."),
        ("Commercial statistics tools", "Use Python/SciPy or JASP."),
        ("Commercial diagram/design tools", "Use draw.io, Flutter, Matplotlib and free design assets with verified licences."),
        ("New test devices", "Use Android Emulator and existing/borrowed university-approved devices."),
        ("Paid survey platform", "Use university-approved forms or privacy-reviewed free forms with anonymous IDs."),
    ], font_size=8.7)

    doc.add_heading("5. Marks-Focused Evidence Checklist", level=1)
    add_table(doc, ["Assessment strength", "Evidence to show examiner"], [
        ("Originality", "Integrated Sinhala on-device HCR plus Grade 1/2 Sinhala readability checking and educator review."),
        ("Methodology", "DSRM traceability, frozen splits, baselines, ablation, expert labels and explicit evaluation gates."),
        ("Technical depth", "CNN/classifier comparison, calibration, latency, integration tests, failure handling and security rules."),
        ("Critical thinking", "Honest limitations, error analysis, alternative-method comparison and threats to validity."),
        ("Reproducibility", "Notebook, environment, hashes, Model Card, Data Card, source tag and README."),
        ("Evaluation", "Model metrics separated from software tests, usability and learning outcomes."),
        ("Ethics", "Data minimisation, consent, child privacy, adult oversight and retention controls."),
        ("Communication", "Clear diagrams, numbered tables, IEEE references, one RQ and objective-to-evidence conclusions."),
    ], font_size=8.7)

    doc.add_heading("6. Final Definition of Done", level=1)
    add_callout(doc, "The project is ready for final submission only when", "The exact shipped models are reproducible; five-letter CNN mapping passes; the text classifier is evaluated against baselines on independently labelled data; visible workflows contain no false scoring or placeholders; tests and privacy controls pass; human evidence has proper approval; the Results and Discussion are traceable to archived outputs; and the final APK demonstrates the core research without a paid service.", PALE_TEAL)


def supervisor_record_and_refs(doc):
    doc.add_heading("7. Supervisor Stage Approval Record", level=1)
    add_body(doc, "Use this table during weekly meetings. A stage should move to 'Approved/Complete' only after its outputs are shown.")
    add_table(doc, ["Stage", "Evidence shown", "Supervisor decision", "Next action / date"], [
        ("0–1", "", "", ""), ("2", "", "", ""), ("3", "", "", ""),
        ("4", "", "", ""), ("5", "", "", ""), ("6–7", "", "", ""),
        ("8", "", "", ""), ("9", "", "", ""), ("10–11", "", "", ""),
    ], font_size=8.5)
    add_body(doc, "Supervisor name: ______________________________   Signature: ____________________   Date: ____________")
    add_body(doc, "Student signature: ______________________________   Date: ____________________")

    doc.add_heading("References and Official Free-Tier Sources", level=1)
    refs = [
        '[1] Flutter, “Build apps for any screen,” 2026. [Online]. Available: https://flutter.dev/. [Accessed: Aug. 21, 2026].',
        '[2] Google Research, “Google Colab Frequently Asked Questions,” 2026. [Online]. Available: https://research.google.com/colaboratory/faq.html. [Accessed: Aug. 21, 2026].',
        '[3] Firebase, “Firebase pricing,” 2026. [Online]. Available: https://firebase.google.com/pricing. [Accessed: Aug. 21, 2026].',
        '[4] Firebase, “Firebase pricing plans,” 2026. [Online]. Available: https://firebase.google.com/docs/projects/billing/firebase-pricing-plans. [Accessed: Aug. 21, 2026].',
        '[5] Android Developers, “Install Android Studio,” 2026. [Online]. Available: https://developer.android.com/studio/install.html. [Accessed: Aug. 21, 2026].',
        '[6] Google Developers, “ML pipelines,” 2025. [Online]. Available: https://developers.google.com/machine-learning/managing-ml-projects/pipelines.',
        '[7] UNICEF Innocenti, “Guidance on AI and Children, Version 3.0,” Dec. 2025. [Online]. Available: https://www.unicef.org/innocenti/reports/policy-guidance-ai-children.',
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Pt(20)
        p.paragraph_format.first_line_indent = Pt(-20)
        p.paragraph_format.space_after = Pt(5)


def build():
    doc = setup_document()
    doc.sections[0].header.paragraphs[0].text = "ReadBuddy AI — Zero-Cost Development Roadmap"
    doc.sections[0].header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title = "ReadBuddy AI Zero-Cost Stage-by-Stage Development Roadmap"
    doc.core_properties.author = "L. M. I. Silva"
    doc.core_properties.subject = "Final Year Research Development Plan — Student ID 28277"
    title_page(doc)
    front_matter(doc)
    overview(doc)
    stages_0_to_3(doc)
    stages_4_to_7(doc)
    stages_8_to_11(doc)
    schedule_and_control(doc)
    supervisor_record_and_refs(doc)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
