from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "documents"
OUTPUT_PATH = OUTPUT_DIR / "ReadBuddy_AI_Viva_Questions_and_Model_Answers.docx"

PURPLE = "5B45F5"
TEAL = "13AAA2"
NAVY = "1C2340"
PALE_PURPLE = "F1EFFF"
PALE_TEAL = "EAF9F7"
PALE_YELLOW = "FFF6D8"
WHITE = "FFFFFF"
GREY = "65708A"


def set_cell_fill(cell, colour):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), colour)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(text)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_callout(doc, title, body, colour=PALE_PURPLE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_fill(cell, colour)
    set_cell_margins(cell, top=160, start=180, bottom=160, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    return table


def add_qa(doc, number, question, answer, note=None):
    table = doc.add_table(rows=2 + (1 if note else 0), cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"

    q_cell = table.cell(0, 0)
    set_cell_fill(q_cell, PURPLE)
    set_cell_margins(q_cell)
    q = q_cell.paragraphs[0]
    q.paragraph_format.space_after = Pt(0)
    run = q.add_run(f"Q{number}. {question}")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(WHITE)

    a_cell = table.cell(1, 0)
    set_cell_fill(a_cell, "FFFFFF")
    set_cell_margins(a_cell)
    a = a_cell.paragraphs[0]
    a.paragraph_format.space_after = Pt(0)
    label = a.add_run("Model answer: ")
    label.bold = True
    label.font.color.rgb = RGBColor.from_string(TEAL)
    a.add_run(answer)

    if note:
        n_cell = table.cell(2, 0)
        set_cell_fill(n_cell, PALE_YELLOW)
        set_cell_margins(n_cell)
        n = n_cell.paragraphs[0]
        n.paragraph_format.space_after = Pt(0)
        label = n.add_run("සිංහල මතක සටහන: ")
        label.bold = True
        n.add_run(note)

    doc.add_paragraph().paragraph_format.space_after = Pt(1)


SECTIONS = [
    (
        "A. Research Foundation",
        [
            (
                "What problem does your research solve?",
                "My research addresses the lack of child-friendly, Sinhala-focused digital support for early literacy. ReadBuddy AI helps Grade 1 and Grade 2 learners practise letters, tracing, vocabulary, reading and graded tasks, while giving teachers evidence-based progress information.",
                "ගැටලුව = මුල් ශ්‍රේණි දරුවන්ට personalised Sinhala literacy support සහ ගුරුවරුන්ට එකම තැනක progress evidence නොමැති වීම.",
            ),
            (
                "Why did you select this research topic?",
                "Early literacy requires frequent practice and individual feedback, but a teacher may not be able to observe every attempt. I selected this topic to investigate how mobile learning and lightweight AI can provide immediate practice feedback and support teacher decision-making.",
                "AI ගුරුවරයා replace කරන්නේ නැහැ; repeated practice සහ feedback support කරනවා.",
            ),
            (
                "What is the main aim of the project?",
                "The main aim is to design and evaluate an intelligent bilingual mobile learning system that supports Grade 1 and Grade 2 literacy development through adaptive activities, Sinhala handwriting recognition and progress-based coaching.",
                None,
            ),
            (
                "What is the main research contribution?",
                "The contribution is the integration of Sinhala handwritten-letter recognition, trace-shape verification, Grade 1/2 text-difficulty classification, structured learning activities and teacher-facing progress analytics in one child-friendly application.",
                "Contribution කියන්නේ app screens විතරක් නොවෙයි—AI, verification, pedagogy සහ analytics එකට integrate කිරීම.",
            ),
            (
                "Why did you limit the scope to Grades 1 and 2?",
                "The scope was limited to Grades 1 and 2 to keep the curriculum, dataset and evaluation manageable and reliable. The available Grade 3 evidence was insufficient for a defensible model, so extending to Grade 3 is identified as future work.",
                "Scope limitation එක weakness එකක් ලෙස නොව scientific boundary එකක් ලෙස explain කරන්න.",
            ),
            (
                "How is this research rather than only a mobile application?",
                "The project defines a research problem, develops and evaluates computational components, measures model performance on held-out data, studies design decisions, and produces evidence about how AI-assisted feedback can support early literacy. The mobile application is the implementation and evaluation platform.",
                "Research evidence: dataset, model, experiments, metrics, evaluation සහ limitations.",
            ),
        ],
    ),
    (
        "B. AI and Machine-Learning Components",
        [
            (
                "How many trained or classification models are included in the final system?",
                "The final system contains two principal classification components: the custom CNN for Sinhala handwritten-letter recognition and the Grade 1/2 text-difficulty classifier. The Learning Coach is a progress-based decision and recommendation component, not a separately trained model.",
                "Correct answer: models/components දෙකයි. Learning Coach එක trained model එකක් කියන්න එපා.",
            ),
            (
                "Is there an AI chatbot in the final system?",
                "No. The final user-facing design uses a controlled AI Learning Coach instead of an open-ended chatbot. This keeps guidance curriculum-focused, predictable, safer for children and directly connected to recorded learning evidence.",
                "Codebase එකේ legacy Gemini-related code තිබුණත් final feature claim එක chatbot නොව Learning Coach.",
            ),
            (
                "Why did you use a CNN for handwritten-letter recognition?",
                "CNNs are suitable for image classification because they learn spatial patterns such as edges, curves, stroke arrangements and local shapes. These characteristics are important when distinguishing visually similar Sinhala handwritten letters.",
                None,
            ),
            (
                "What exactly does the CNN do?",
                "The learner's drawing is converted into a standardised image representation and passed to the TensorFlow Lite CNN. The model returns a predicted class and confidence score, which are then used together with verification logic to assess the attempt.",
                None,
            ),
            (
                "What is the reported CNN test accuracy?",
                "The custom CNN achieved 88.61% accuracy on the held-out test dataset. I describe this specifically as test accuracy only because it was calculated on a test split that was isolated from model training and tuning.",
                "88.61% කියන්න. Validation accuracy සහ test accuracy mix කරන්න එපා.",
            ),
            (
                "What is the difference between accuracy and confidence?",
                "Accuracy is an aggregate measure showing the percentage of correct predictions across a labelled evaluation dataset. Confidence is the model's score for one particular prediction; high confidence does not automatically mean that the prediction is correct.",
                None,
            ),
            (
                "Why is CNN confidence alone insufficient for tracing assessment?",
                "A classifier may assign high confidence to an incorrect or poorly formed drawing if it contains familiar visual features. Therefore, the application also checks geometric similarity and trace evidence before showing a successful result.",
                "මේක screenshot වල තිබුණු false-positive problem එකට scientific answer එක.",
            ),
            (
                "How does the hybrid tracing evaluation work?",
                "The system standardises the drawing, obtains CNN identity evidence and calculates shape similarity against the expected form. It then applies verification thresholds and consistency rules to produce the final score and child-friendly feedback.",
                "CNN = letter identity; shape check = written form similarity.",
            ),
            (
                "Why does the CNN have 454 output classes?",
                "The output space follows the structure of the training dataset and its encoded glyph classes. The application maps those model classes to the smaller curriculum-facing set of verified letters shown to Grade 1 and Grade 2 learners.",
                "454 class කියන්නේ UI එකේ letters 454ක් තියෙනවා කියන එක නොවෙයි.",
            ),
            (
                "What does the text-difficulty classifier do?",
                "It analyses characteristics of reading content and assigns an appropriate Grade 1 or Grade 2 difficulty level. This supports age-appropriate content selection and prevents learners from receiving material that is unnecessarily difficult.",
                None,
            ),
            (
                "Is the Learning Coach a trained machine-learning model?",
                "No. It is a transparent decision-support component that analyses task attempts, recent activity, strengths and weak areas to generate prioritised recommendations. Its rules are intentionally controlled so that recommendations are explainable to teachers and parents.",
                "Coach ගැන ‘trained AI model’ කියලා overclaim කරන්න එපා.",
            ),
        ],
    ),
    (
        "C. Dataset, Training and Evaluation",
        [
            (
                "How did you prepare the image data?",
                "The images were validated, assigned to verified classes, standardised to the input format required by the CNN, and separated into training, validation and held-out test subsets. Invalid or ambiguous samples were excluded where necessary.",
                None,
            ),
            (
                "Why are training, validation and test sets separated?",
                "The training set is used to learn model parameters, the validation set supports model selection and tuning, and the held-out test set provides a final estimate of generalisation to unseen samples.",
                None,
            ),
            (
                "How did you reduce the risk of data leakage?",
                "I kept the held-out test set separate from training and tuning, controlled class mappings, and used split manifests and data-quality checks. Related or excluded samples were handled so that evaluation would not be artificially inflated.",
                None,
            ),
            (
                "Which evaluation metrics did you use?",
                "I used overall accuracy and class-level precision, recall and F1-score, together with the classification report and training-history evidence. A confusion matrix is also useful for identifying visually similar classes that the model confuses.",
                None,
            ),
            (
                "If test accuracy is 88.61%, what explains the remaining errors?",
                "Errors can result from large handwriting variation, visually similar Sinhala forms, limited examples for particular classes, incomplete strokes, image preprocessing differences and low-quality drawings. The verification layer and retry feedback reduce the practical impact of these errors.",
                None,
            ),
            (
                "How do you know the result is reproducible?",
                "The project retains the training pipeline, model card, architecture description, run manifest, metrics, classification report and the shipped TensorFlow Lite model. These artefacts document how the reported result was produced.",
                "Repository එකේ training/stage3_cnn_pipeline.py, metrics JSON සහ model card evidence තියෙනවා.",
            ),
            (
                "Why is a high score not automatically proof of perfect handwriting?",
                "A score reflects the implemented similarity and classification criteria; it is not a complete pedagogical judgement. The system is a practice aid, while a teacher remains responsible for evaluating stroke order, fine motor control and curriculum correctness.",
                None,
            ),
        ],
    ),
    (
        "D. System Design and Features",
        [
            (
                "What are the main features of ReadBuddy AI?",
                "The system includes Grade 1 and Grade 2 learning paths, Sinhala and English learning content, letters and pillam, tracing with AI-assisted checking, vocabulary, reading stories, levelled tasks and quizzes, progress tracking, a Learning Coach, profiles and a teacher dashboard.",
                None,
            ),
            (
                "Why are Grade 1 and Grade 2 learning paths different?",
                "Grade 1 focuses on letter recognition, sound association, basic tracing and simple vocabulary. Grade 2 progresses to pillam, word construction, reading comprehension and more demanding tasks, following a simple-to-advanced learning sequence.",
                None,
            ),
            (
                "Why did you use child-friendly images and colours?",
                "Young learners depend strongly on visual recognition and require large, simple interaction targets. Images, colour coding, limited text and consistent navigation reduce cognitive load and make practice more engaging.",
                None,
            ),
            (
                "How does the language design support both Sinhala- and English-medium learners?",
                "The interface language and learning language are treated as separate choices. A learner may use an English interface while studying Sinhala letters, or use a Sinhala interface while practising English content.",
                "UI language ≠ learning language කියන distinction එක මතක තියාගන්න.",
            ),
            (
                "How does a teacher monitor an individual learner?",
                "The teacher can link a student profile and review activity history, tracing attempts, quiz and task performance, reading evidence, weak areas, inactivity indicators and coaching recommendations. Grade-based grouping helps the teacher locate learners quickly.",
                None,
            ),
            (
                "Why is daily activity synchronisation important?",
                "A final average alone cannot show whether a child is practising consistently. Daily synchronisation helps teachers distinguish improvement, inactivity, repeated difficulty and missing practice, enabling timely intervention.",
                None,
            ),
            (
                "Why did you use Flutter?",
                "Flutter provides a single maintainable codebase, strong Android support, rapid child-friendly UI development and suitable packages for local TensorFlow Lite inference and Firebase integration.",
                None,
            ),
            (
                "Why did you use Firebase?",
                "Firebase supports authentication, cloud data storage and synchronisation between learner and teacher views. It allows progress evidence to be updated without embedding sensitive service credentials in the application.",
                None,
            ),
            (
                "Which functions can work offline and which require internet access?",
                "The bundled learning content, local exercises, TensorFlow Lite inference and some local progress operations can work on-device. Authentication, cross-device synchronisation and teacher access to newly uploaded progress require connectivity.",
                None,
            ),
        ],
    ),
    (
        "E. Testing, Ethics, Limitations and Future Work",
        [
            (
                "How did you test the complete application?",
                "I used static analysis, automated unit and widget tests, model evaluation on held-out data, and manual scenario testing on an Android emulator or device. Scenarios covered navigation, bilingual content, tracing, tasks, progress updates and teacher–student data flow.",
                None,
            ),
            (
                "How do you protect children's data?",
                "The design follows data minimisation, authenticated access and role-based separation of learner and teacher views. Only learning-related information needed for progress support should be stored, and production deployment must use appropriate Firebase security rules and consent procedures.",
                None,
            ),
            (
                "What ethical concerns did you consider?",
                "The main concerns are child privacy, misleading AI feedback, bias across handwriting styles, safe language, transparency and over-reliance on automated scores. The system presents AI as support and preserves teacher oversight.",
                None,
            ),
            (
                "What are the main limitations?",
                "The main limitations are the size and representativeness of the child-handwriting data, possible confusion between similar letter forms, limited long-term school deployment, device and input variation, and the current restriction to Grades 1 and 2.",
                None,
            ),
            (
                "What would you improve in future work?",
                "Future work should collect a larger consented child-handwriting dataset, conduct longitudinal classroom evaluation, improve class balancing and calibration, add more curriculum grades, strengthen speech-based reading assessment and test personalised recommendations through controlled studies.",
                None,
            ),
            (
                "Can this system replace a teacher?",
                "No. ReadBuddy AI provides practice, immediate feedback and organised evidence. A teacher is still required for pedagogy, emotional support, diagnosis of complex learning needs and final educational judgement.",
                "මෙය අනිවාර්යයෙන්ම ‘No’ කියලා start කරන්න.",
            ),
            (
                "What would you demonstrate if the examiner asks for a live demo?",
                "I would show the Grade 1/2 path selection, one learning activity, a valid and invalid tracing attempt, a task or quiz, the updated learner progress, and the corresponding teacher view or Learning Coach recommendation. I would also explain which result is produced by the CNN and which is produced by rule-based analytics.",
                None,
            ),
        ],
    ),
]


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, colour in (
        ("Title", 26, NAVY),
        ("Subtitle", 12, GREY),
        ("Heading 1", 17, PURPLE),
        ("Heading 2", 13, TEAL),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(colour)
        style.font.bold = style_name != "Subtitle"

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.add_run("ReadBuddy AI — Viva Preparation  |  ")
    add_page_number(footer_p)

    # Cover page
    doc.add_paragraph().paragraph_format.space_after = Pt(30)
    logo_line = doc.add_paragraph()
    logo_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = logo_line.add_run("📚  READBUDDY AI")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Viva Questions and\nModel Answers")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Grade 1–2 Intelligent Reading and Writing Support System")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    add_callout(
        doc,
        "Purpose of this document",
        "A concise defence guide for explaining the research problem, system architecture, AI components, model evaluation, educational rationale, limitations and future work.",
        PALE_TEAL,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(15)
    facts = doc.add_table(rows=5, cols=2)
    facts.alignment = WD_TABLE_ALIGNMENT.CENTER
    facts.style = "Table Grid"
    fact_rows = [
        ("Final scope", "Grade 1 and Grade 2 literacy support"),
        ("Classification components", "2 — CNN letter recognition + Grade 1/2 text difficulty"),
        ("CNN held-out test accuracy", "88.61%"),
        ("User-facing intelligent support", "Progress-based AI Learning Coach"),
        ("Important clarification", "No open-ended chatbot in the final system"),
    ]
    for i, (label, value) in enumerate(fact_rows):
        left, right = facts.rows[i].cells
        set_cell_fill(left, PALE_PURPLE)
        set_cell_fill(right, "FFFFFF")
        set_cell_margins(left)
        set_cell_margins(right)
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        left.paragraphs[0].add_run(label).bold = True
        right.paragraphs[0].add_run(value)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Prepared for final project viva and demonstration")
    r.italic = True
    r.font.color.rgb = RGBColor.from_string(GREY)

    doc.add_section(WD_SECTION.NEW_PAGE)

    # Quick defence map
    add_heading(doc, "Quick Defence Map", 1)
    add_callout(
        doc,
        "30-second project introduction",
        "ReadBuddy AI is a bilingual mobile learning system for Grade 1 and Grade 2 children. It combines curriculum-based literacy activities with a custom CNN for Sinhala handwritten-letter recognition, a Grade 1/2 text-difficulty classifier, hybrid tracing verification, progress analytics and a teacher-facing dashboard. Its purpose is to provide immediate practice feedback and actionable learning evidence while keeping teachers in control.",
        PALE_TEAL,
    )

    add_heading(doc, "The four points to repeat consistently", 2)
    add_bullet(doc, "Scope: Grade 1 and Grade 2 only.", "Scope:")
    add_bullet(doc, "Models: Two principal classification components.", "Models:")
    add_bullet(doc, "Accuracy: The custom CNN achieved 88.61% held-out test accuracy.", "Accuracy:")
    add_bullet(doc, "Coach: The Learning Coach is an explainable progress-based recommendation component, not an open-ended chatbot or separately trained model.", "Coach:")

    add_heading(doc, "Simple architecture explanation", 2)
    arch = doc.add_table(rows=5, cols=3)
    arch.alignment = WD_TABLE_ALIGNMENT.CENTER
    arch.style = "Table Grid"
    headers = ("Component", "Purpose", "Type")
    for j, h in enumerate(headers):
        cell = arch.cell(0, j)
        set_cell_fill(cell, PURPLE)
        set_cell_margins(cell)
        rr = cell.paragraphs[0].add_run(h)
        rr.bold = True
        rr.font.color.rgb = RGBColor.from_string(WHITE)
    set_repeat_table_header(arch.rows[0])
    architecture_rows = [
        ("Custom CNN", "Recognises Sinhala handwritten letter identity", "Trained image classifier"),
        ("Text-difficulty classifier", "Assigns Grade 1/2 reading difficulty", "Classification component"),
        ("Shape verification", "Checks geometric similarity of tracing", "Deterministic verification"),
        ("Learning Coach", "Finds weak areas and suggests next actions", "Rule/progress-based analytics"),
    ]
    for i, row in enumerate(architecture_rows, start=1):
        for j, value in enumerate(row):
            cell = arch.cell(i, j)
            set_cell_margins(cell)
            set_cell_fill(cell, "FFFFFF" if i % 2 else PALE_TEAL)
            cell.paragraphs[0].add_run(value)

    doc.add_section(WD_SECTION.NEW_PAGE)

    number = 1
    for section_title, questions in SECTIONS:
        add_heading(doc, section_title, 1)
        for question, answer, note in questions:
            add_qa(doc, number, question, answer, note)
            number += 1

    add_heading(doc, "Final Viva Checklist", 1)
    add_callout(
        doc,
        "Avoid these incorrect claims",
        "Do not say that the final system contains three trained models. Do not call confidence the same as accuracy. Do not describe validation accuracy as test accuracy. Do not claim that AI replaces the teacher. Do not present the final system as an open-ended chatbot.",
        PALE_YELLOW,
    )
    add_heading(doc, "Before entering the viva", 2)
    checklist = [
        "Memorise the research problem, aim, four objectives and final Grade 1/2 scope.",
        "Know the difference between the CNN, shape verification, text classifier and Learning Coach.",
        "Memorise the CNN held-out test accuracy: 88.61%.",
        "Prepare one correct tracing and one incorrect tracing for the demonstration.",
        "Prepare a student profile with visible task, tracing and progress evidence.",
        "Be ready to explain one limitation honestly and one realistic future improvement.",
        "If a live feature fails, explain the architecture and show retained test evidence or screenshots calmly.",
    ]
    for item in checklist:
        add_bullet(doc, item)

    add_heading(doc, "Recommended closing statement", 2)
    add_callout(
        doc,
        "Closing answer",
        "This research demonstrates that a lightweight, child-friendly mobile system can combine on-device Sinhala handwriting recognition with structured literacy activities and explainable progress analytics. The present results are promising for Grade 1 and Grade 2 support, while larger child-centred datasets and longitudinal classroom studies are required before wider deployment.",
        PALE_TEAL,
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document())
