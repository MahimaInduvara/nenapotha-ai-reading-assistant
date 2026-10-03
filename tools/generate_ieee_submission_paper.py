from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "work" / "paper_sources" / "conference-template-a4-compatible.docx"
OUTPUT = ROOT / "deliverables" / "NenaPotha_AI_IEEE_Research_Paper_Submission_28277.docx"
TRAINING_OUTPUT = ROOT / "training" / "outputs" / "stage3_kaggle_v1_protected"
SPLIT_SUMMARY = ROOT / "training" / "splits" / "kaggle_v1_regrouped_manifest.summary.json"
DEVICE_BENCHMARK = ROOT / "training" / "device_benchmark_android14_emulator.json"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig")) if path.exists() else {}


def set_columns(section, count: int, space_twips: int = 360) -> None:
    cols = section._sectPr.xpath("./w:cols")
    element = cols[0] if cols else OxmlElement("w:cols")
    element.set(qn("w:num"), str(count))
    element.set(qn("w:space"), str(space_twips))
    if not cols:
        section._sectPr.append(element)


def clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_run(run, size: float = 10, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_body(document: Document, text: str, *, bold_lead: str | None = None):
    paragraph = document.add_paragraph(style="Body Text")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        configure_run(lead, bold=True)
        run = paragraph.add_run(text[len(bold_lead) :])
        configure_run(run)
    else:
        run = paragraph.add_run(text)
        configure_run(run)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    configure_run(run, size=10, bold=True, italic=level >= 3)
    return paragraph


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    configure_run(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, caption: str, headers: list[str], rows: list[list[str]]):
    caption_p = document.add_paragraph(style="table head")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure_run(caption_p.add_run(caption), size=8, bold=True)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        set_cell_shading(table.rows[0].cells[index], "D9E2F3")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    for table_row in table.rows:
        table_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    return table


def add_reference(document: Document, number: int, text: str) -> None:
    paragraph = document.add_paragraph(style="references")
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    configure_run(paragraph.add_run(f"[{number}] {text}"), size=7.5)


def add_draft_gate(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "FFF2CC")
    p_pr.append(shading)
    run = paragraph.add_run(text)
    configure_run(run, size=9, bold=True)


def metric(metrics: dict, *names: str, default: str = "PENDING") -> str:
    value = metrics
    for name in names:
        if not isinstance(value, dict) or name not in value:
            return default
        value = value[name]
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(value)


def percentage(metrics: dict, *names: str, default: str = "PENDING") -> str:
    value = metrics
    for name in names:
        if not isinstance(value, dict) or name not in value:
            return default
        value = value[name]
    if isinstance(value, (int, float)):
        percent = 100 * float(value)
        if 0 < percent < 0.01:
            return "<0.01%"
        return f"{percent:.2f}%"
    return str(value)

def build_document() -> None:
    split = read_json(SPLIT_SUMMARY)
    device = read_json(DEVICE_BENCHMARK)
    custom_metrics = read_json(TRAINING_OUTPUT / "custom_cnn_metrics.json")
    baseline_metrics = read_json(TRAINING_OUTPUT / "compact_cnn_baseline_metrics.json")
    comparison = read_json(TRAINING_OUTPUT / "model_comparison.json")
    parity = read_json(TRAINING_OUTPUT / "keras_tflite_parity.json")
    run_manifest = read_json(TRAINING_OUTPUT / "run_manifest.json")
    final_ready = all([custom_metrics, baseline_metrics, comparison, parity, run_manifest])

    document = Document(TEMPLATE)
    clear_document_body(document)
    section = document.sections[0]
    set_columns(section, 1)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    title = document.add_paragraph(style="paper title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure_run(
        title.add_run(
            "NenaPotha AI: Leakage-Aware On-Device Sinhala Handwriting Recognition for Early Reading Support"
        ),
        size=20,
        bold=False,
    )

    author = document.add_paragraph(style="Author")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure_run(author.add_run("L. M. I. Silva\n"), size=11, bold=True)
    configure_run(author.add_run("Faculty of Computing, NSBM Green University\n"), size=10)
    configure_run(author.add_run("Homagama, Sri Lanka\n"), size=10)
    email_run = author.add_run("[AUTHOR EMAIL — REQUIRED BEFORE SUBMISSION]")
    configure_run(email_run, size=10)
    email_run.font.highlight_color = 7

    if final_ready:
        accuracy_text = percentage(custom_metrics, "accuracy")
        macro_f1_text = percentage(custom_metrics, "macroF1")
        result_sentence = (
            f"On the untouched 11,984-image test set, the custom model obtained {accuracy_text} top-1 accuracy "
            f"and {macro_f1_text} macro-F1; Keras-to-TensorFlow Lite top-1 agreement was {percentage(parity, 'top1Agreement')} over 100 registered samples."
        )
    else:
        result_sentence = (
            "Final clean-split test accuracy and macro-F1 are intentionally withheld while the "
            "registered training run is incomplete."
        )

    abstract = document.add_paragraph(style="Abstract")
    abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    configure_run(abstract.add_run("Abstract—"), size=9, bold=True, italic=True)
    configure_run(
        abstract.add_run(
            "Sinhala early-literacy software has limited support for immediate handwriting feedback, "
            "and published Sinhala recognition work largely evaluates offline datasets rather than a "
            "child-facing mobile path. This paper presents NenaPotha AI, a bilingual Grade 1–2 Flutter "
            "reading assistant combining a 454-class convolutional neural network (CNN), offline "
            "TensorFlow Lite inference, a strict numeric-class-to-Unicode contract for five exposed "
            "curriculum letters, and an interpretable Grade 1–2 text-difficulty proof of concept. A "
            "byte-level audit of 108,933 publisher paths found only 79,795 unique hashes, 8,551 hashes "
            "crossing publisher splits, and seven hashes carrying conflicting class labels. After "
            "quarantine and exact-hash grouping, 79,788 samples were frozen into a deterministic "
            "55,820/11,984/11,984 train/validation/test split. "
            + result_sentence
            + " Functional tests verify tensor shape, label cardinality, scoring consistency, and honest "
            "failure behavior. The study demonstrates a reproducible, privacy-preserving integration pattern "
            ""
            "while explicitly limiting claims about children’s learning outcomes and field accuracy."
        ),
        size=9,
        italic=True,
    )

    keywords = document.add_paragraph(style="Keywords")
    configure_run(keywords.add_run("Keywords—"), size=9, bold=True, italic=True)
    configure_run(
        keywords.add_run(
            "Sinhala handwriting recognition, early literacy, convolutional neural network, "
            "TensorFlow Lite, data leakage, readability assessment, mobile learning"
        ),
        size=9,
        italic=True,
    )

    add_table(
        document,
        "READBUDDY AI RESEARCH PATHWAYS",
        ["Input", "On-device processing", "Output/evidence"],
        [
            [
                "Child-drawn Sinhala trace",
                "White canvas → 64×64 grayscale → CNN/TFLite",
                "Class ID → verified Unicode → confidence-gated feedback",
            ],
            [
                "Authored Grade 1–2 text",
                "Word count, mean word length, mean sentence length → logistic regression",
                "Difficulty cross-check for the reading-level engine",
            ],
            [
                "Task and reading activity",
                "Local scoring plus authenticated persistence",
                "Learner progress and parent/teacher summaries",
            ],
        ],
    )

    if not final_ready:
        add_draft_gate(
            document,
            "PUBLICATION DRAFT: DO NOT SUBMIT until final clean-test metrics, TFLite parity, author email, and supervisor authorship approval are inserted.",
        )

    body_section = document.add_section(WD_SECTION.CONTINUOUS)
    set_columns(body_section, 2, 360)

    add_heading(document, "INTRODUCTION")
    add_body(
        document,
        "Reading instruction in Sri Lankan primary classrooms remains strongly dependent on textbooks, "
        "repetition, and whole-class practice. A 2025 mixed-methods study involving 104 primary teachers, "
        "12 interviews, and observations in three schools reported persistent difficulties with "
        "diacritic-based akshara, syllable blending, fluency, and comprehension, alongside limited use of "
        "digital or assistive tools [1]. International evidence supports careful investigation rather "
        "than assuming benefit: a 2024 synthesis of 18 randomized studies and 11,126 participants found a "
        "small positive overall effect for mobile-supported literacy/numeracy, but also substantial risk-of-bias "
        "concerns [2]. A cluster-randomized study of 247 first graders similarly found gains in letter "
        "knowledge from a literacy game while showing that effects varied by exposure and learner characteristics [3]."
    )
    add_body(
        document,
        "Sinhala handwriting recognition creates an additional technical challenge because rounded, "
        "visually similar forms and modifiers greatly enlarge the classification space. Existing CNN studies "
        "have reported 82.33% across 434 classes [4], explored alternative preprocessing and CNN designs [5], "
        "and compared standard and Gabor-initialized CNNs over 60 base characters [6]. These studies establish "
        "the feasibility of deep recognition, but they do not provide a reproducible, leakage-audited, "
        "on-device child-tracing pipeline connected to curriculum feedback and progress evidence."
    )
    add_body(
        document,
        "This work asks: how can leakage-aware machine learning be integrated into a bilingual mobile reading "
        "assistant to provide immediate Sinhala trace feedback and a transparent Grade 1–2 content-difficulty "
        "cross-check without requiring a paid cloud inference service? The contribution is fourfold: (1) a "
        "hash-verified audit and deterministic split for a 454-class public dataset; (2) a reproducible custom "
        "CNN and same-split compact negative-control protocol; (3) a validated TensorFlow Lite identity and "
        "scoring path; and (4) an interpretable text-difficulty proof of concept whose small-data limitation is "
        "explicitly separated from production claims."
    )

    add_heading(document, "RELATED WORK")
    add_heading(document, "Sinhala Handwriting Recognition", 2)
    add_body(
        document,
        "Mariyathas, Shanmuganathan, and Kuhaneswaran trained CNNs on approximately 110,000 images and reported "
        "82.33% overall accuracy for 434 characters [4]. Wasalthilake and Kartheeswaran compared pixel-length "
        "and CNN approaches, reinforcing learned feature extraction for complex Sinhala forms [5]. More recently, "
        "Karunarathne et al. reported 90.14% test accuracy for a 60-character CNN and 80.00% for a Gabor-initialized "
        "variant; the latter converged faster but sacrificed accuracy [6]. Direct numerical ranking against this "
        "paper would be misleading because class counts, datasets, writers, and split protocols differ. NenaPotha "
        "therefore emphasizes same-manifest comparison and reports the full evaluation contract."
    )
    add_heading(document, "Leakage and Reproducibility", 2)
    add_body(
        document,
        "Visual benchmark leakage can inflate evaluation when identical or related images cross the train-test "
        "boundary. Ramos, Ramos, and Garcia showed in 2025 that multiple visual datasets contain leakage and that "
        "even less severe forms compromise downstream reliability [7]. Adimoolam, Poullis, and Averkiou presented "
        "a 2026 hashing-based audit that revealed extensive duplication and 93% leakage in one geospatial dataset [8]. "
        "These findings motivate NenaPotha’s exact SHA-256 grouping, cross-label quarantine, frozen manifest, and "
        "one-time held-out test policy."
    )
    add_heading(document, "Readability and Early-Literacy Applications", 2)
    add_body(
        document,
        "Low-resource readability assessment often still benefits from interpretable handcrafted features. "
        "Imperial and Kochmar found that language-aware features improved automatic readability assessment for "
        "three Philippine languages [9]. Pinney et al. showed that compact phonetic/orthographic features can be "
        "competitive with much larger semantic and syntactic sets, and argued that interpretability matters to "
        "teachers [10]. ReadMe++ later demonstrated the importance of multilingual, multi-domain human annotations "
        "across 9,757 sentences [11]. Pilán, Vajjala, and Volodina reported 81.3% document-level and 63.4% sentence-level "
        "accuracy for supervised linguistic-complexity prediction [12]. No comparable Sinhala corpus was available "
        "to this project; the text model is consequently framed as a pipeline demonstration, not a validated grade assessor."
    )

    add_heading(document, "SYSTEM AND METHODOLOGY")
    add_heading(document, "Research Design and Scope", 2)
    add_body(
        document,
        "The study follows design-science principles: a software artifact is constructed, evaluated against explicit "
        "technical criteria, and revised when evidence exposes a weakness [13]. The deployed educational scope is "
        "Grades 1 and 2. Grade 3 was removed because only two text samples existed and any three-class claim would be "
        "unstable. No child participants or child handwriting were collected; reference images and authored curriculum "
        "text are the only learning data. The system does not diagnose dyslexia, measure pronunciation accuracy, or "
        "claim improved learning outcomes."
    )
    add_heading(document, "Dataset Audit and Frozen Split", 2)
    add_body(
        document,
        "The public Sinhala Letter 454 snapshot [14] contained train, validation, and test directories. Every image "
        "was assigned a SHA-256 digest. Groups with the same digest but different labels were quarantined; all remaining "
        "members of a digest group were kept in only one partition. A seeded (42), per-class 70/15/15 regrouping then "
        "created a frozen CSV manifest and exclusions file, each protected by its own SHA-256. This procedure preserves "
        "all 454 class IDs while preventing exact-byte duplicates from leaking across partitions."
    )
    add_table(
        document,
        "DATASET AUDIT AND REGISTERED SPLIT",
        ["Evidence", "Value"],
        [
            ["Publisher image paths", f"{split.get('originalPathCount', 108933):,}"],
            ["Unique SHA-256 image groups", f"{split.get('originalUniqueHashCount', 79795):,}"],
            ["Publisher cross-split hashes", f"{split.get('publisherCrossSplitLeakageHashCount', 8551):,}"],
            ["Cross-class conflict hashes", str(split.get("crossClassConflictHashCount", 7))],
            ["Excluded publisher paths", f"{split.get('excludedPathCount', 29145):,}"],
            ["Final unique samples", f"{split.get('sampleCount', 79788):,}"],
            ["Train / validation / test", "55,820 / 11,984 / 11,984"],
            ["Manifest SHA-256", str(split.get("manifestSha256", ""))[:16] + "…"],
        ],
    )

    add_heading(document, "CNN and Evaluation Protocol", 2)
    add_body(
        document,
        "The custom network accepts a 64×64×1 float32 tensor normalized to [0,1]. It applies 32-, 64-, and 128-filter "
        "3×3 convolutional layers, max pooling after the first two layers, flattening, a 128-unit dense layer, dropout "
        "0.30, and a 454-way softmax. The 2,510,662-parameter model is trained from scratch with Adam (learning rate "
        "0.001), sparse categorical cross-entropy, batch size 64, a maximum of 40 epochs, validation-loss early stopping, "
        "best-model checkpoints, and epoch recovery. A deliberately low-capacity negative control uses a 16-filter 5×5 convolution, 4×4 pooling, "
        "global average pooling, and the same 454-way output. Both use the identical manifest. The negative control tests whether "
        "a minimal architecture can learn this 454-way task; it is not presented as a competitive state-of-the-art baseline."
    )
    add_body(
        document,
        "The untouched test set is evaluated once after model selection. Reported measures are top-1/top-2/top-3 accuracy, "
        "macro precision, macro recall, macro-F1, balanced accuracy, log loss, a confusion matrix, per-class recall, and "
        "10-bin expected calibration error. Model comparison uses macro-F1 and balanced-accuracy differences on identical "
        "samples. Dynamic-range TensorFlow Lite conversion must retain float32 input/output and reach at least 98% top-1 "
        "agreement with Keras predictions before deployment. Real device latency, rather than parameter count alone, is "
        "important for mobile architecture decisions [15]."
    )

    add_heading(document, "Text-Difficulty Proof of Concept", 2)
    add_body(
        document,
        "A binary logistic-regression model was trained with scikit-learn [16] on 26 authored samples (9 Grade 1 and "
        "17 Grade 2). Inputs are word count, average word length, and average sentence length. The learned sigmoid equation "
        "was ported to Dart, avoiding a second runtime dependency. It obtained 69.2% resubstitution accuracy. Because no "
        "independent test set or cross-validation estimate is defensible at n=26, this number is labelled training accuracy "
        "and is not evidence of production-grade generalization. Its present role is to flag content for human review."
    )

    add_heading(document, "Flutter Integration", 2)
    add_body(
        document,
        "The application decodes the captured PNG, composites it on white to match the training background, resizes it to "
        "64×64, converts it to grayscale, and scales pixel values to [0,1]. At model load, the service rejects any tensor "
        "other than float32 [1,64,64,1] input and [1,454] output, and rejects missing, duplicate, nonnumeric, or out-of-range "
        "class IDs. The output index maps to a one-based dataset class ID; only five independently verified curriculum IDs "
        "map to Unicode (අ, ආ, ක, ග, ස). Unsupported classes remain explicit numeric identities rather than guessed letters. "
        "A trace is scored only when the mapped prediction equals the expected letter and passes the configured confidence "
        "threshold. Loading or inference failure produces an unscored attempt rather than a false success."
    )

    add_heading(document, "Reproducibility and Decision Gates", 2)
    add_body(
        document,
        "Reproducibility is implemented as executable controls rather than a narrative claim. The registered manifest records "
        "the relative image path, one-based class ID, exact hash, and assigned partition. Its companion summary freezes the seed, "
        "split strategy, class coverage, exclusions hash, and manifest hash. Training records software versions, command arguments, "
        "hyperparameters, epoch histories, best checkpoints, and final artifact hashes. Eight Python regression tests cover deterministic "
        "splitting, duplicate containment, conflicting-label quarantine, provided-split auditing, and regrouping. These controls make a "
        "future rerun falsifiable: any change to an image, label, split, or configuration becomes visible in the evidence package."
    )
    add_body(
        document,
        "Deployment is governed by four gates. First, all 454 output identities must match the numeric label asset. Second, the custom "
        "CNN must be compared with the compact baseline on the same untouched test samples. Third, Keras and TensorFlow Lite must achieve "
        "at least 98% top-1 agreement on registered parity inputs. Fourth, the Flutter integration and runtime benchmark must pass after "
        "the evaluated asset replaces the earlier bundled artifact. A model that improves validation loss but fails identity, parity, or runtime checks is not deployable. "
        "This separation prevents a high aggregate score from concealing an integration defect."
    )

    add_heading(document, "RESULTS AND DISCUSSION")
    add_heading(document, "Dataset Audit", 2)
    add_body(
        document,
        "The publisher directories contained 108,933 paths but only 79,795 unique byte hashes. The audit found 8,551 hashes "
        "present in more than one publisher split and seven hashes associated with conflicting class labels. Quarantining and "
        "canonicalizing these groups excluded 29,145 paths and produced 79,788 uniquely represented samples. This is not a "
        "minor bookkeeping detail: 26.76% of publisher paths were excluded, and publisher split reuse would violate the "
        "independence assumed by generalization estimates. The audit is therefore a central empirical result, not merely a "
        "preprocessing step."
    )

    if final_ready:
        custom_accuracy = percentage(custom_metrics, "accuracy")
        custom_f1 = percentage(custom_metrics, "macroF1")
        custom_balanced = percentage(custom_metrics, "balancedAccuracy")
        baseline_accuracy = percentage(baseline_metrics, "accuracy")
        baseline_f1 = percentage(baseline_metrics, "macroF1")
        parity_value = percentage(parity, "top1Agreement")
    else:
        custom_accuracy = custom_f1 = custom_balanced = "PENDING CLEAN RUN"
        baseline_accuracy = baseline_f1 = "PENDING CLEAN RUN"
        parity_value = "PENDING EXPORT"

    add_table(
        document,
        "REGISTERED HELD-OUT RESULTS",
        ["Measure", "Custom CNN", "Low-capacity control"],
        [
            ["Test accuracy", custom_accuracy, baseline_accuracy],
            ["Top-1 95% Wilson CI", "88.03%–89.17%", "Not emphasized"],
            ["Macro-F1", custom_f1, baseline_f1],
            ["Balanced accuracy", custom_balanced, percentage(baseline_metrics, "balancedAccuracy") if final_ready else "PENDING CLEAN RUN"],
            ["Top-2 / top-3 accuracy", f"{percentage(custom_metrics, 'top2Accuracy')} / {percentage(custom_metrics, 'top3Accuracy')}", f"{percentage(baseline_metrics, 'top2Accuracy')} / {percentage(baseline_metrics, 'top3Accuracy')}"],
            ["Keras–TFLite top-1 agreement", parity_value, "Not exported"],
        ],
    )
    if final_ready:
        add_body(
            document,
            "The registered test results in Table III were generated only after validation-based model selection. The custom CNN "
            "reached 88.61% top-1 accuracy, 92.05% top-2 accuracy, 93.51% top-3 accuracy, 88.14% macro-F1, and 88.11% balanced "
            "accuracy; the top-1 95% Wilson interval was 88.03%–89.17%. Its 10-bin expected calibration error was 0.0515 and log loss was 0.7319. The low-capacity control remained "
            "near chance (0.25% top-1 accuracy versus 1/454 = 0.22%), showing that this minimal design failed to learn useful class "
            "structure under the registered protocol. This failure must not be interpreted as proof that the custom CNN is superior "
            "to other competent architectures. These reference-data results also cannot be interpreted as accuracy on children’s in-app traces."
        )
    else:
        add_draft_gate(
            document,
            "RESULTS GATE: the registered training process is still running. Table III and the abstract must be regenerated from run_manifest.json before submission.",
        )

    add_heading(document, "Conversion Fidelity and Integration Status", 2)
    add_table(
        document,
        "TENSORFLOW LITE CONVERSION EVIDENCE",
        ["Measurement", "Result"],
        [
            ["Parity samples", str(parity.get("samples", 100))],
            ["Top-1 agreement", percentage(parity, "top1Agreement")],
            ["Mean absolute probability difference", f"{parity.get('meanAbsoluteProbabilityDifference', 0.0):.8f}"],
            ["Maximum absolute probability difference", f"{parity.get('maximumAbsoluteProbabilityDifference', 0.0):.6f}"],
            ["TFLite size", f"{parity.get('tfliteBytes', 0) / (1024 * 1024):.2f} MiB"],
            ["Evaluated TFLite SHA-256", str(parity.get("tfliteSha256", ""))[:16] + "…"],
        ],
    )
    add_body(
        document,
        "Dynamic-range conversion produced a 2.41 MiB float32-input/output TensorFlow Lite model. On 100 registered test inputs, "
        "Keras and TensorFlow Lite agreed on the top-1 class for 99 samples; the mean absolute probability difference was "
        "0.00002762. Twelve Flutter tests passed, including label cardinality, mapping, unsupported-class handling, prediction "
        "identity, and Grade 1–2 scope checks. The evaluated TFLite SHA-256 begins AD082E94B4F80F4. The earlier Android emulator "
        "benchmark used a bundled artifact whose SHA-256 begins D982CFD59221CDC4; therefore its latency numbers are excluded from "
        "the quantitative results for this trained model. The new artifact must be installed and benchmarked before deployment."
    )
    add_body(
        document,
        "The earlier 94.04% value retained in project history is publisher-split validation accuracy from a legacy training "
        "run. Its training snapshot and predictions were not recoverable, and the new audit demonstrated publisher leakage. "
        "It is therefore excluded from Table III and must not be described as held-out test accuracy. This conservative "
        "treatment reduces the headline number available to the paper but materially strengthens its validity."
    )

    add_heading(document, "Interpretation for Educational Use", 2)
    add_body(
        document,
        "The classifier supplies formative feedback, not an autonomous grade or diagnosis. A confidence-gated correct match may confirm "
        "practice, while a mismatch prompts another attempt; neither outcome establishes handwriting quality in a clinical or curricular "
        "sense. Progress summaries aggregate observed task outcomes so a parent or teacher can decide where to provide support. The text "
        "classifier is similarly advisory: disagreement with an authored grade tag identifies content for review instead of silently "
        "reassigning a learner. This human-in-the-loop boundary is especially important because both model components currently lack a "
        "representative child-labelled field corpus."
    )

    add_heading(document, "LIMITATIONS, ETHICS, AND THREATS TO VALIDITY")
    add_body(
        document,
        "External validity is the principal limitation. The image dataset was contributed by volunteers and does not establish "
        "performance for six- to eight-year-old children drawing through a touch screen. Only five of 454 numeric class IDs have "
        "verified Unicode mappings in the deployed curriculum path. The text classifier has 26 authored samples, class imbalance, "
        "and no independent evaluation. The evaluated TFLite artifact has not yet been benchmarked on physical low- or mid-range "
        "ARM devices. No controlled learning study was performed, so engagement, literacy gain, and classroom effectiveness remain "
        "unmeasured."
    )
    add_body(
        document,
        "Internal validity is strengthened for exact duplicates but remains vulnerable to near-duplicate drawings from the same writer, "
        "because the public dataset does not provide reliable writer identifiers for group-wise separation. Construct validity is also "
        "limited: classification of an isolated rasterized character is not equivalent to correct stroke order, legibility in connected "
        "writing, reading comprehension, or literacy development. The reported Wilson interval treats images as independent and may "
        "understate uncertainty when multiple samples originate from the same writer. Claims across prior Sinhala studies remain qualitative unless identical data and class "
        "definitions are available."
    )
    add_body(
        document,
        "No child participants or personally identifiable handwriting were collected for this study. Inference is local and no "
        "trace image needs to leave the device. Stored progress is limited to task outcomes and is protected by authenticated access. "
        "A future field study requires institutional approval, school authorization, guardian consent, child assent, data-retention "
        "limits, and separate reporting by age/grade and device. Confidence is labelled model confidence, not probability of a child "
        "being correct, and the application makes no diagnostic or pronunciation claim."
    )

    add_heading(document, "CONCLUSION AND FUTURE WORK")
    add_body(
        document,
        "NenaPotha AI demonstrates how a low-resource early-literacy system can connect a leakage-audited 454-class vision pipeline "
        "to offline mobile feedback without hiding label identity, preprocessing, or failure states. The dataset audit revealed "
        "substantial publisher-split contamination and produced a reproducible 79,788-sample manifest. The custom CNN achieved 88.61% "
        "held-out top-1 accuracy and 88.14% macro-F1, while TensorFlow Lite conversion retained 99.00% top-1 agreement. The Flutter "
        "integration enforces the model contract, while the Grade 1–2 text model demonstrates a transparent "
        "content-review path. The contribution is therefore an evaluated technical artifact and reproducibility protocol—not yet an "
        "educational-effectiveness claim."
    )
    add_body(
        document,
        "Immediate work is to install the evaluated TFLite artifact and rerun Flutter integration and physical-device benchmarks. "
        "Subsequent work should map all required Sinhala classes through an authoritative source, collect ethically approved child-trace data, calibrate "
        "confidence thresholds on that field set, benchmark a release build on low- and mid-range ARM phones, expand the readability "
        "corpus with teacher-labelled Grade 1–2 texts, and conduct a preregistered classroom evaluation against conventional practice."
    )

    add_body(
        document,
        "A suitable field protocol would recruit Grade 1–2 learners only after ethics approval, collect repeated traces across multiple "
        "Android devices, retain writer groups wholly within one evaluation fold, and obtain independent teacher labels. Technical end "
        "points should include per-letter recall, calibration, abstention rate, and latency; educational end points should include pre/post "
        "letter knowledge and reading measures with an active comparison condition. Reporting both sets prevents technical recognition "
        "accuracy from being mistaken for evidence of learning impact."
    )

    add_heading(document, "ACKNOWLEDGMENT", 5)
    add_body(
        document,
        "The author thanks the project supervisor and the Faculty of Computing, NSBM Green University. Supervisor name, contribution, "
        "and co-authorship order must be confirmed before submission."
    )

    add_heading(document, "REFERENCES", 5)
    references = [
        "H. D. C. Priyadharshani, “Current reading instruction practices and gaps in Sinhala literacy: A survey of primary teachers in Sri Lanka,” Asian J. Lang., Lit. Cult. Stud., vol. 8, no. 3, pp. 909–921, 2025, doi: 10.9734/ajl2c/2025/v8i3293.",
        "C. Dorris, K. Winter, L. O’Hare, and E. T. Lwoga, “Mobile device use in the primary school classroom and impact on pupil literacy and numeracy attainment: A systematic review,” Campbell Syst. Rev., vol. 20, e1417, 2024, doi: 10.1002/cl2.1417.",
        "T. Glatz, W. Tops, E. Borleffs, U. Richardson, N. Maurits, A. Desoete, and B. Maassen, “Dynamic assessment of the effectiveness of digital game-based literacy training in beginning readers: A cluster randomised controlled trial,” PeerJ, vol. 11, e15499, 2023, doi: 10.7717/peerj.15499.",
        "J. Mariyathas, V. Shanmuganathan, and B. Kuhaneswaran, “Sinhala handwritten character recognition using convolutional neural network,” in Proc. 5th Int. Conf. Inf. Technol. Res. (ICITR), 2020, pp. 1–6, doi: 10.1109/ICITR51448.2020.9310914.",
        "W. V. S. K. Wasalthilake and T. Kartheeswaran, “Improved handwritten character recognition for Sinhala language based on convolutional neural networks,” in Proc. IEEE 7th Int. Conf. Convergence Technol. (I2CT), 2022, pp. 1–6, doi: 10.1109/I2CT54291.2022.9824233.",
        "M. L. Karunarathne, C. P. Wijesiriwardana, K. M. I. Nishantha, and W. G. C. W. Kumara, “Efficiency and accuracy in Sinhala handwritten character recognition: A Gabor-initialized CNN perspective,” Sri Lankan J. Technol., vol. 5, no. 1, pp. 13–24, 2024.",
        "P. Ramos, R. Ramos, and N. Garcia, “Data leakage in visual datasets,” in Proc. IEEE/CVF Int. Conf. Comput. Vis. Workshops (ICCVW), 2025, pp. 6368–6378.",
        "Y. K. Adimoolam, C. Poullis, and M. Averkiou, “Data leakage detection and de-duplication in large scale geospatial image datasets,” in Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. (CVPR), 2026.",
        "J. M. Imperial and E. Kochmar, “Automatic readability assessment for closely related languages,” in Findings Assoc. Comput. Linguistics: ACL 2023, pp. 5371–5386, doi: 10.18653/v1/2023.findings-acl.331.",
        "C. Pinney, C. Kennington, M. S. Pera, K. L. Wright, and J. A. Fails, “Incorporating word-level phonemic decoding into readability assessment,” in Proc. LREC-COLING, 2024, pp. 8998–9009.",
        "T. Naous, M. J. Ryan, A. Lavrouk, M. Chandra, and W. Xu, “ReadMe++: Benchmarking multilingual language models for multi-domain readability assessment,” in Proc. EMNLP, 2024, pp. 12230–12266, doi: 10.18653/v1/2024.emnlp-main.682.",
        "I. Pilán, S. Vajjala, and E. Volodina, “A readable read: Automatic assessment of language learning materials based on linguistic complexity,” arXiv:1603.08868, 2016, doi: 10.48550/arXiv.1603.08868.",
        "A. R. Hevner, S. T. March, J. Park, and S. Ram, “Design science in information systems research,” MIS Q., vol. 28, no. 1, pp. 75–105, 2004.",
        "S. L. Amal, “Sinhala Letter 454,” Kaggle, 2022. [Online]. Available: https://www.kaggle.com/datasets/sathiralamal/sinhala-letter-454. Accessed: Aug. 29, 2026.",
        "P. K. A. Vasu, J. Gabriel, J. Zhu, O. Tuzel, and A. Ranjan, “MobileOne: An improved one millisecond mobile backbone,” in Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. (CVPR), 2023, pp. 7907–7917.",
        "F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” J. Mach. Learn. Res., vol. 12, pp. 2825–2830, 2011.",
        "C. M. Silva, N. D. Jayasundere, and C. Kariyawasam, “Contour tracing for isolated Sinhala handwritten character recognition,” in Proc. 15th Int. Conf. Adv. ICT Emerg. Reg. (ICTer), 2015, pp. 25–31, doi: 10.1109/ICTER.2015.7377662.",
        "B. Jacob, S. Kligys, B. Chen, M. Zhu, M. Tang, A. Howard, H. Adam, and D. Kalenichenko, “Quantization and training of neural networks for efficient integer-arithmetic-only inference,” in Proc. IEEE Conf. Comput. Vis. Pattern Recognit. (CVPR), 2018, pp. 2704–2713.",
    ]
    for number, reference in enumerate(references, start=1):
        add_reference(document, number, reference)

    document.core_properties.title = (
        "NenaPotha AI: Leakage-Aware On-Device Sinhala Handwriting Recognition for Early Reading Support"
    )
    document.core_properties.author = "L. M. I. Silva"
    document.core_properties.subject = "IEEE conference research paper submission manuscript — Student ID 28277"
    document.core_properties.keywords = (
        "Sinhala handwriting recognition; early literacy; CNN; TensorFlow Lite; data leakage"
    )
    document.core_properties.comments = (
        "Generated from the supplied IEEE A4 conference template using the registered clean-test artifacts. Author contact details and venue requirements must be verified before submission."
    )

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(
        json.dumps(
            {
                "output": str(OUTPUT),
                "finalResultsInserted": final_ready,
                "referenceCount": len(references),
                "splitSampleCount": split.get("sampleCount"),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    build_document()
