from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "documents"
OUTPUT_PATH = OUTPUT_DIR / "ReadBuddy_Supervisory_Meeting_Minutes_Draft.docx"

STUDENT_NAME = "L.M.I. Silva (Mahima)"
STUDENT_ID = "28277"
SUPERVISOR = "Mr. Anton Jayakody"
PROJECT_TITLE = (
    "Design and Development of an AI-Based Interactive Reading and Speaking "
    "Assistant for Primary School Students"
)

MEETINGS = [
    {
        "number": "01",
        "date": "26/02/2026",
        "discussed": [
            "Introduced the proposed research problem concerning limited personalised support for early-grade reading, writing and speaking practice.",
            "Discussed the intended beneficiaries: Grade 1 and Grade 2 learners, parents and primary-school teachers.",
            "Reviewed the initial project aim and possible features, including letter learning, tracing, vocabulary, reading practice and progress monitoring.",
            "Considered the feasibility of developing the solution as a Flutter-based Android application.",
        ],
        "next": [
            "Conduct a focused literature review on early literacy, Sinhala handwriting recognition and mobile learning.",
            "Refine the research problem, research question, aim and measurable objectives.",
            "Prepare the initial functional requirements, scope and system concept.",
        ],
    },
    {
        "number": "02",
        "date": "01/04/2026",
        "discussed": [
            "Reviewed preliminary literature and identified educational, technical and progress-monitoring gaps.",
            "Discussed limiting the final research scope to Grade 1 and Grade 2 to maintain curriculum and evaluation reliability.",
            "Reviewed the proposed architecture: Flutter application, Firebase services, on-device AI inference and teacher-facing progress views.",
            "Discussed child-friendly bilingual design, data privacy and the need to keep learning-language selection separate from interface language.",
        ],
        "next": [
            "Develop the initial Chapter 1, Chapter 2 and Chapter 3 drafts.",
            "Prepare low-fidelity UI wireframes and the first application prototype.",
            "Define the Sinhala letter dataset, curriculum content and model-evaluation plan.",
        ],
    },
    {
        "number": "03",
        "date": "19/06/2026",
        "discussed": [
            "Reviewed progress on Chapters 1–3, including the research gap, objectives, literature review and methodology.",
            "Confirmed the use of a custom CNN for Sinhala handwritten-letter recognition.",
            "Discussed a separate Grade 1/2 text-difficulty classifier and the need to present it honestly as a small proof-of-concept component.",
            "Reviewed dataset preparation, training/validation/test separation, class mapping and evidence required to prevent data leakage.",
        ],
        "next": [
            "Revise Chapters 1–3 according to the feedback received.",
            "Clean and preprocess the image dataset and establish reproducible data splits.",
            "Train the CNN baseline and prepare Grade 1/2 text samples for the difficulty classifier.",
            "Continue implementing the Grade 1/2 learning prototype.",
        ],
    },
    {
        "number": "04",
        "date": "03/07/2026",
        "discussed": [
            "Demonstrated the early mobile prototype, including bilingual navigation, grade-based content and letter-tracing interaction.",
            "Reviewed the CNN preprocessing contract: 64×64 grayscale input, normalisation and TensorFlow Lite deployment.",
            "Discussed the risk of using CNN confidence alone to assess handwriting and the need for an additional geometric shape check.",
            "Agreed on evaluation evidence: accuracy, precision, recall, F1-score, classification report and confusion analysis.",
        ],
        "next": [
            "Complete the controlled CNN training and held-out test evaluation.",
            "Convert and integrate the selected CNN model into the Flutter application.",
            "Implement hybrid tracing verification using CNN identity and shape-similarity evidence.",
            "Add automated tests for label mapping and tracing-result logic.",
        ],
    },
    {
        "number": "05",
        "date": "30/07/2026",
        "discussed": [
            "Reviewed the custom CNN result of 88.61% accuracy on the held-out test dataset and the retained training evidence.",
            "Reviewed the 454-output TensorFlow Lite model, verified class mapping and on-device inference integration.",
            "Discussed incorrect tracing acceptance and the combined CNN-plus-shape verification solution.",
            "Reviewed the Grade 1/2 logistic-regression text classifier trained on 26 content samples, with 69.2% training accuracy.",
            "Confirmed that validation accuracy, test accuracy and prediction confidence must be reported using correct terminology.",
        ],
        "next": [
            "Complete CNN and shape-verification integration for the supported Sinhala letters.",
            "Improve tracing feedback and test both correct and incorrect drawings.",
            "Document the text classifier as a proof of concept and state its dataset limitation.",
            "Retain model cards, metrics, architecture and reproducibility evidence for the thesis appendices.",
        ],
    },
    {
        "number": "06",
        "date": "18/08/2026",
        "discussed": [
            "Reviewed the integrated Grade 1/2 learning paths, stories, quizzes, vocabulary and tracing features.",
            "Discussed replacing open-ended chatbot behaviour with a controlled, progress-based Learning Coach suitable for children.",
            "Reviewed the teacher dashboard, student linking, daily activity synchronisation, weak-area analysis and coaching records.",
            "Reviewed voice-reading comparison, bilingual content behaviour and the separation of interface language from learning language.",
            "Identified the need for complete system testing and evidence-backed Results and Discussion chapters.",
        ],
        "next": [
            "Complete teacher/student progress synchronisation and verify all major user flows.",
            "Conduct static analysis, automated tests and representative emulator/device testing.",
            "Capture clear application screenshots and prepare figures for the thesis.",
            "Draft the Results, Discussion, Conclusion and Future Work chapters using accurate model evidence and IEEE references.",
        ],
    },
    {
        "number": "07",
        "date": "30/09/2026",
        "discussed": [
            "Demonstrated the near-final ReadBuddy AI application and reviewed the complete Grade 1/2 learner and teacher workflows.",
            "Reviewed the two classification components: CNN letter recognition and Grade 1/2 text-difficulty classification.",
            "Reviewed the hybrid tracing check, voice-reading comparison, Learning Coach and explainable teacher analytics.",
            "Demonstrated the teacher-facing Text Difficulty Checker and clarified that its result is advisory and requires teacher review.",
            "Reviewed final thesis structure, screenshots, research-component explanation, limitations and viva preparation.",
        ],
        "next": [
            "Complete final thesis formatting, IEEE reference checking, appendices, tables and figures.",
            "Run final Flutter analysis and the complete automated test suite; retain the result evidence.",
            "Prepare a reliable live demonstration and backup screenshots for the viva.",
            "Submit the final thesis and project artefacts after supervisor review and approval.",
        ],
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, value=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_label_value(doc, label, value):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(12.5)
    left, right = table.rows[0].cells
    left.width = Cm(4.2)
    right.width = Cm(12.5)
    for cell in (left, right):
        set_cell_margins(cell, 40)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    left.paragraphs[0].add_run(label).bold = True
    right.paragraphs[0].add_run(value)
    return table


def add_minutes_box(doc, heading, items):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(heading)
    r.bold = True
    r.font.size = Pt(10.5)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.height = Cm(5.4)
    set_cell_margins(cell, 140)
    for index, item in enumerate(items):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.style = doc.styles["List Bullet"]
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.add_run(item)
    return table


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(30, 30, 30)

    for meeting_index, meeting in enumerate(MEETINGS):
        if meeting_index:
            doc.add_page_break()

        brand = doc.add_paragraph()
        brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
        brand.paragraph_format.space_after = Pt(1)
        run = brand.add_run("NSBM GREEN UNIVERSITY TOWN")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(36, 104, 62)

        title_table = doc.add_table(rows=1, cols=2)
        title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        title_table.autofit = False
        title_table.columns[0].width = Cm(13.8)
        title_table.columns[1].width = Cm(3.0)
        title_cell, meeting_cell = title_table.rows[0].cells
        title_cell.width = Cm(13.8)
        meeting_cell.width = Cm(3.0)
        set_cell_margins(title_cell, 80)
        set_cell_margins(meeting_cell, 80)
        title_p = title_cell.paragraphs[0]
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run("Final Year Project – Supervisory Meeting Minutes")
        title_run.bold = True
        title_run.font.size = Pt(15)
        meeting_cell.paragraphs[0].add_run(
            f"Meeting No: {meeting['number']}"
        ).bold = True
        meeting_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        note = doc.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.paragraph_format.space_after = Pt(7)
        note_run = note.add_run("DRAFT FOR SUPERVISOR REVIEW")
        note_run.bold = True
        note_run.font.size = Pt(8.5)
        note_run.font.color.rgb = RGBColor(180, 65, 50)

        add_label_value(doc, "Date", meeting["date"])
        add_label_value(doc, "Project Title", PROJECT_TITLE)
        add_label_value(doc, "Name of the Student", STUDENT_NAME)
        add_label_value(doc, "Student ID", STUDENT_ID)
        add_label_value(doc, "Name of the Supervisor", SUPERVISOR)

        add_minutes_box(doc, "Items discussed:", meeting["discussed"])
        add_minutes_box(
            doc,
            "Items to be completed before the next supervisory meeting:",
            meeting["next"],
        )

        signature = doc.add_paragraph()
        signature.paragraph_format.space_before = Pt(12)
        signature.paragraph_format.space_after = Pt(1)
        signature.add_run("............................................................")
        label = doc.add_paragraph("Supervisor (Signature & Date)")
        label.paragraph_format.space_after = Pt(0)
        label.runs[0].font.size = Pt(9.5)

        footer_note = doc.add_paragraph()
        footer_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_note.paragraph_format.space_before = Pt(5)
        footer_run = footer_note.add_run(
            "Supervisor signature must be completed by the supervisor after "
            "reviewing and approving the recorded minutes."
        )
        footer_run.italic = True
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = RGBColor(90, 90, 90)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document())
