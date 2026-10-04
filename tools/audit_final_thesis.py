from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


WORD_RE = re.compile(r"\b[\w%–—’'-]+\b", re.UNICODE)


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document._body)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document._body)


def block_text(block) -> str:
    if isinstance(block, Paragraph):
        return block.text.strip()
    return " ".join(
        paragraph.text.strip()
        for row in block.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        if paragraph.text.strip()
    )


def count_words(text: str) -> int:
    return len(WORD_RE.findall(text))


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: audit_final_thesis.py THESIS.docx")

    path = Path(sys.argv[1])
    document = Document(path)
    blocks = list(iter_blocks(document))

    body_parts: list[str] = []
    abstract_parts: list[str] = []
    in_body = False
    in_abstract = False
    for block in blocks:
        text = block_text(block)
        if text == "ABSTRACT":
            in_abstract = True
            continue
        if text == "TABLE OF CONTENTS":
            in_abstract = False
        if text == "1 INTRODUCTION":
            in_body = True
        if text == "REFERENCES":
            in_body = False
        if in_body:
            body_parts.append(text)
        if in_abstract:
            abstract_parts.append(text)

    paragraph_text = "\n".join(p.text for p in document.paragraphs)
    appendix_headings = [
        p.text.strip()
        for p in document.paragraphs
        if p.style.name == "Heading 1"
        and re.match(r"^APPENDIX [A-T]:", p.text.strip())
    ]
    figure_captions = [
        p.text.strip()
        for p in document.paragraphs
        if p.style.name == "Figure Caption"
        and re.match(r"^Figure (?:[1-6]|[A-T])\.", p.text.strip())
    ]
    screenshot_captions = [
        p.text.strip()
        for p in document.paragraphs
        if p.style.name == "Figure Caption"
        and re.match(r"^Figure I\.\d+:", p.text.strip())
    ]
    references = [
        p.text.strip()
        for p in document.paragraphs
        if re.match(r"^\[\d+\]", p.text.strip())
    ]
    forbidden = [
        token
        for token in (
            "114/114",
            "114 automated",
            "Additional clean screenshots",
            "ReadBuddy AI",
            "Update the field to generate",
        )
        if token in paragraph_text
    ]

    print(f"file={path}")
    print(f"body_words={count_words(' '.join(body_parts))}")
    print(f"abstract_words={count_words(' '.join(abstract_parts))}")
    print(f"appendix_count={len(appendix_headings)}")
    print(f"appendix_range={appendix_headings[0]} -> {appendix_headings[-1]}")
    print(f"figure_caption_count={len(figure_captions)}")
    print(f"screenshot_caption_count={len(screenshot_captions)}")
    print(f"inline_picture_count={len(document.inline_shapes)}")
    print(f"table_count={len(document.tables)}")
    print(f"reference_count={len(references)}")
    print(f"forbidden_tokens={forbidden}")

    errors = []
    if not 10_000 <= count_words(" ".join(body_parts)) <= 12_000:
        errors.append("Body word count is outside 10,000–12,000")
    if not 200 <= count_words(" ".join(abstract_parts)) <= 300:
        errors.append("Abstract word count is outside 200–300")
    if len(appendix_headings) != 20:
        errors.append("Appendix A–T set is incomplete")
    if len(screenshot_captions) != 14:
        errors.append("Not all 14 screenshot figures are captioned")
    if len(references) < 20:
        errors.append("Reference list is unexpectedly short")
    if forbidden:
        errors.append("Outdated/placeholder terms remain")
    if errors:
        raise SystemExit("AUDIT FAILED: " + "; ".join(errors))
    print("audit=PASS")


if __name__ == "__main__":
    main()
